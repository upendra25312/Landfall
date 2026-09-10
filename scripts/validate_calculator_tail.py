"""Inspect the four calculator tail adapters and retain their actual Excel export."""
import argparse
import asyncio
import json
import re
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src/calc"))
from adapters import ADAPTERS
from smoke import SAMPLE, CALCULATOR_URL
from driver import _APPLY_MODULE, _CONFIGURED, _add_with_retry, _clear_estimate
from configuration import validate_configuration

SERVICES = ("azure-bastion", "application-gateway", "azure-monitor", "load-balancer")
SNAPSHOT = """idx => { const m=document.querySelectorAll('.row.product-module')[idx];
return {text:m.innerText, controls:[...m.querySelectorAll('input[name],select[name]')].map(e=>({
name:e.name, value:e.value, type:e.type, visible:!!e.getClientRects().length,
options:e.options?[...e.options].map(o=>({value:o.value,text:o.text})):undefined}))}; }"""


async def run(output, region="sweden-central", variant=False):
    from playwright.async_api import async_playwright
    output.mkdir(parents=True, exist_ok=True)
    observations = []
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True)
        page = await browser.new_page(viewport={"width": 1440, "height": 1200})
        page.set_default_timeout(60000)
        await page.goto(CALCULATOR_URL, wait_until="domcontentloaded")
        await page.wait_for_selector("button.pickerItem__button--add, .service-info-picker-button")
        await _clear_estimate(page)
        for service in SERVICES:
            ad = ADAPTERS[service]
            config = dict(SAMPLE[service])
            if variant:
                config.update({
                    "azure-bastion": {"tier": "premium", "hours": 100, "scale_units": 5, "outbound_data_gb": 20},
                    "application-gateway": {"tier": "wafv2", "hours": 100, "capacity_units": 3, "data_processed_gb": 2048},
                    "azure-monitor": {"log_data_ingestion_gb": 180, "basic_logs_gb": 30, "interactive_retention_months": 6},
                    "load-balancer": {"rules": 11, "data_processed_gb": 2048},
                }[service])
            before = await page.evaluate(f"() => {_CONFIGURED}.length")
            await _add_with_retry(page, ad["product"], before)
            await page.wait_for_timeout(1500)
            idx = await page.evaluate("""() => { const a=[...document.querySelectorAll('.row.product-module')];
                return a.findLastIndex(m=>m.querySelector('input[name],select[name]')); }""")
            initial = await page.evaluate(SNAPSHOT, idx)
            fields = [[name, value, kind] for name, value, kind in ad["fields"](config)]
            result = await page.evaluate(_APPLY_MODULE, {"idx": idx, "region": region, "fields": fields,
                                                       "regionOptional": service in ("bandwidth", "azure-dns")})
            await page.wait_for_timeout(1500)
            observations.append({"service": service, "config": config, "fields": fields,
                                 "initial": initial, "apply": result, "final": await page.evaluate(SNAPSHOT, idx)})
            try:
                observations[-1]["assumptions"] = validate_configuration(service, config, result)
                observations[-1]["valid"] = True
            except ValueError as exc:
                observations[-1]["valid"] = False
                observations[-1]["error"] = str(exc)
            (output / "controls.json").write_text(json.dumps(observations, indent=2), encoding="utf-8")
            print(service, result, flush=True)
        async with page.expect_download(timeout=90000) as download:
            await page.click("button.export-button")
        await (await download.value).save_as(output / "calculator.xlsx")
        from openpyxl import load_workbook
        workbook = load_workbook(output / "calculator.xlsx", data_only=True)
        rows = {r[1]: r for r in workbook.active.iter_rows(values_only=True) if r[1] in [ADAPTERS[s]["product"] for s in SERVICES]}
        assert len(rows) == len(SERVICES), "calculator export omitted a service"
        for observation in observations:
            row = rows[ADAPTERS[observation["service"]]["product"]]
            cost = re.search(r"Monthly cost\s+\$([\d,.]+)", observation["final"]["text"])
            assert cost and abs(float(cost[1].replace(",", "")) - float(row[5])) < 0.01
            actual = {c["name"]: c["value"] for c in observation["final"]["controls"]}
            for name, expected, kind in observation["fields"]:
                if kind == "accordion" or name not in actual:
                    continue  # missing controls remain explicitly validated above
                assert str(actual[name]) == str(expected), (name, actual[name], expected)
            observation["export"] = {"region": row[3], "description": row[4], "monthly_cost": row[5], "ui_total_matches": True}
        (output / "controls.json").write_text(json.dumps(observations, indent=2), encoding="utf-8")
        await page.screenshot(path=str(output / "calculator.png"))
        await browser.close()
    return 1 if any(not o["valid"] for o in observations) else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--region", default="sweden-central")
    parser.add_argument("--variant", action="store_true")
    parser.add_argument("--services", nargs="+")
    args = parser.parse_args()
    if args.services:
        SERVICES = tuple(args.services)
    raise SystemExit(asyncio.run(run(args.output, args.region, args.variant)))
