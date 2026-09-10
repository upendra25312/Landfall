"""
Resource Plan OpenPyXL Workbook Generator (PRD E15C.9–E15C.14).

Generates a comprehensive, 15-sheet Excel workbook (`resource_plan.xlsx`) using deterministic
openpyxl code without external proprietary templates. Uses native Excel formulas with fullCalcOnLoad.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from resource.model import STANDARD_ROLES

# Palette
_NAVY = "1B365D"
_ACCENT = "0078D4"
_LIGHT_BG = "F3F4F6"
_WARN_BG = "FEF3C7"
_DANGER_BG = "FEE2E2"
_SUCCESS_BG = "D1FAE5"

_HEAD_FILL = PatternFill("solid", fgColor=_NAVY)
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, color=_NAVY, size=14)
_SECTION_FONT = Font(bold=True, color=_NAVY, size=11)
_BOLD_FONT = Font(bold=True, size=10)
_REG_FONT = Font(size=10)
_ITALIC_FONT = Font(italic=True, size=9, color="555555")

_THIN_SIDE = Side(border_style="thin", color="CCCCCC")
_BORDER_ALL = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def _style_header_row(ws, row_idx: int, num_cols: int) -> None:
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER_ALL


def _autofit_columns(ws, min_col: int = 1, max_col: int | None = None, min_width: int = 12, max_width: int = 50) -> None:
    cols = range(min_col, (max_col or ws.max_column) + 1)
    for col_idx in cols:
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            val = row[0].value
            if val is not None:
                # Do not count multiline title rows
                s = str(val)
                if len(s) < 80:
                    max_len = max(max_len, len(s))
        ws.column_dimensions[col_letter].width = min(max_width, max(min_width, max_len + 3))


def generate_resource_workbook(
    demand: dict[str, Any],
    capacity_analysis: dict[str, Any] | None = None,
    commercial: dict[str, Any] | None = None,
    sow_narrative: dict[str, Any] | None = None,
    project_plan: list[dict[str, Any]] | None = None,
) -> bytes:
    """
    Generates a 15-sheet Excel workbook representing the complete resource plan.
    """
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    cap = capacity_analysis or {}
    comm = commercial or {}

    totals = demand.get("totals", {})
    params = demand.get("parameters", {})
    role_demand = demand.get("role_demand", {})
    monthly_sched = demand.get("monthly_schedule", {})
    month_keys = monthly_sched.get("month_keys", [])
    role_month_fte = monthly_sched.get("role_month_fte", {})

    # --------------------------------------------------------------------------
    # Sheet 01: Executive Summary
    # --------------------------------------------------------------------------
    ws01 = wb.active
    ws01.title = "01 Executive Summary"
    ws01["A1"] = "Landfall Migration Resource Plan — Executive Summary"
    ws01["A1"].font = _TITLE_FONT
    ws01["A2"] = "Authoritative deterministic demand model with role-based planning and capacity separation."
    ws01["A2"].font = _ITALIC_FONT

    # KPI Summary Table
    ws01["A4"] = "Core Metric"
    ws01["B4"] = "Value"
    ws01["C4"] = "Unit"
    ws01["D4"] = "Derivation Basis"
    _style_header_row(ws01, 4, 4)

    kpis = [
        ("Total Migration Effort", totals.get("total_person_days", 0), "Person-Days (EAC)", "Scope (servers/apps) + 6R dispositions + contingency"),
        ("Total Working Hours", totals.get("total_hours", 0), "Hours", f"Total PD × {params.get('hours_per_day', 8)} hours/day"),
        ("Project Duration", totals.get("total_months", 0), "Months", "Dated wave schedule or relative duration"),
        ("Average Team Size", totals.get("avg_fte", 0), "FTE", f"Total PD ÷ ({totals.get('total_months', 1)} mo × {params.get('working_days_per_month', 21)} days)"),
        ("Peak Team Size", totals.get("peak_fte", 0), "FTE", "Maximum concurrent required FTE in single month"),
        ("Peak Period", totals.get("peak_month", "N/A"), "Period", "Month with highest concurrent wave execution"),
        ("Migration Waves", params.get("waves", 0), "Waves", "Calculated move groups packed by affinity"),
        ("Total Commercial Cost", comm.get("total_commercial_cost", 0), comm.get("currency", "USD"), "Role PD × configured day rates across delivery models"),
        ("Capacity Status", cap.get("status", "CAPACITY NOT PROVIDED"), "Flag", cap.get("notice", "Pre-sales role-based default")),
    ]

    for r_i, (k, v, u, b) in enumerate(kpis, start=5):
        ws01.cell(row=r_i, column=1, value=k).font = _BOLD_FONT
        cell_v = ws01.cell(row=r_i, column=2, value=v)
        cell_v.alignment = Alignment(horizontal="right")
        cell_v.font = _REG_FONT
        if k == "Total Commercial Cost":
            cell_v.number_format = "#,##0"
        ws01.cell(row=r_i, column=3, value=u).font = _REG_FONT
        ws01.cell(row=r_i, column=4, value=b).font = _ITALIC_FONT

    _autofit_columns(ws01)

    # --------------------------------------------------------------------------
    # Sheet 02: Resource Assumptions
    # --------------------------------------------------------------------------
    ws02 = wb.create_sheet("02 Resource Assumptions")
    ws02["A1"] = "Migration Resource Planning Assumptions"
    ws02["A1"].font = _TITLE_FONT

    ws02["A3"] = "Parameter"
    ws02["B3"] = "Configured Value"
    ws02["C3"] = "Unit"
    ws02["D3"] = "Category"
    _style_header_row(ws02, 3, 4)

    assumptions_data = [
        ("Hours per Working Day", params.get("hours_per_day", 8), "Hours", "Calendar"),
        ("Working Days per Month", params.get("working_days_per_month", 21), "Days", "Calendar"),
        ("Working Days per Week", 5, "Days", "Calendar"),
        ("Planning Mode", demand.get("planning_mode", "MODE 1 — Role-Based Resource Planning"), "Mode", "Governance"),
        ("Data Quality Confidence", params.get("dq_confidence", "Medium"), "Confidence Level", "Data Quality"),
        ("Programme Management Uplift", "15.0%", "Percent of Delivery", "Overhead"),
        ("Technical Governance Uplift", "10.0%", "Percent of Delivery", "Overhead"),
        ("Contingency Buffer", f"{round((totals.get('contingency_pd', 0) / (totals.get('delivery_subtotal_pd', 1) or 1)) * 100, 1)}%", "Percent of Base", "Risk"),
    ]
    for r_i, row in enumerate(assumptions_data, start=4):
        for c_i, val in enumerate(row, start=1):
            ws02.cell(row=r_i, column=c_i, value=val).font = _REG_FONT
    _autofit_columns(ws02)

    # --------------------------------------------------------------------------
    # Sheet 03: Role Catalogue
    # --------------------------------------------------------------------------
    ws03 = wb.create_sheet("03 Role Catalogue")
    ws03["A1"] = "Functional Migration Roles Catalogue (MEG Aligned)"
    ws03["A1"].font = _TITLE_FONT

    headers03 = ["Role ID", "Functional Role Title", "Category", "Default Delivery Model", "Primary Skills", "Responsibilities"]
    for c_i, h in enumerate(headers03, start=1):
        ws03.cell(row=3, column=c_i, value=h)
    _style_header_row(ws03, 3, len(headers03))

    for r_i, r_def in enumerate(STANDARD_ROLES, start=4):
        ws03.cell(row=r_i, column=1, value=r_def.id).font = _REG_FONT
        ws03.cell(row=r_i, column=2, value=r_def.title).font = _BOLD_FONT
        ws03.cell(row=r_i, column=3, value=r_def.category).font = _REG_FONT
        ws03.cell(row=r_i, column=4, value=r_def.default_delivery_model.value).font = _REG_FONT
        ws03.cell(row=r_i, column=5, value=", ".join(r_def.skills[:3])).font = _REG_FONT
        ws03.cell(row=r_i, column=6, value=r_def.description).font = _ITALIC_FONT
    _autofit_columns(ws03)

    # --------------------------------------------------------------------------
    # Sheet 04: Activities
    # --------------------------------------------------------------------------
    ws04 = wb.create_sheet("04 Activities")
    ws04["A1"] = "Migration Workstreams & Activities Breakdown"
    ws04["A1"].font = _TITLE_FONT

    headers04 = ["Workstream", "MEG Phase", "Basis / Sizing Logic", "Total Effort (PD)", "Total Hours"]
    for c_i, h in enumerate(headers04, start=1):
        ws04.cell(row=3, column=c_i, value=h)
    _style_header_row(ws04, 3, len(headers04))

    act_row = 4
    for r in demand.get("requirements", []):
        ws04.cell(row=act_row, column=1, value=r["workstream"]).font = _BOLD_FONT
        ws04.cell(row=act_row, column=2, value=r["phase"]).font = _REG_FONT
        ws04.cell(row=act_row, column=3, value=r["notes"]).font = _ITALIC_FONT
        ws04.cell(row=act_row, column=4, value=r["working_days"]).font = _REG_FONT
        ws04.cell(row=act_row, column=5, value=r["effort_hours"]).font = _REG_FONT
        act_row += 1

    # Total row
    ws04.cell(row=act_row, column=1, value="Total").font = _BOLD_FONT
    ws04.cell(row=act_row, column=4, value=f"=SUM(D4:D{act_row-1})").font = _BOLD_FONT
    ws04.cell(row=act_row, column=5, value=f"=SUM(E4:E{act_row-1})").font = _BOLD_FONT
    _autofit_columns(ws04)

    # --------------------------------------------------------------------------
    # Sheet 05: Resource Demand
    # --------------------------------------------------------------------------
    ws05 = wb.create_sheet("05 Resource Demand")
    ws05["A1"] = "Authoritative Resource Demand Requirements Table"
    ws05["A1"].font = _TITLE_FONT

    headers05 = ["Workstream", "Phase", "Wave", "Activity", "Role", "Delivery Model", "Working Days (PD)", "Hours", "Required FTE"]
    for c_i, h in enumerate(headers05, start=1):
        ws05.cell(row=3, column=c_i, value=h)
    _style_header_row(ws05, 3, len(headers05))

    req_r = 4
    for row in demand.get("requirements", []):
        ws05.cell(row=req_r, column=1, value=row["workstream"]).font = _REG_FONT
        ws05.cell(row=req_r, column=2, value=row["phase"]).font = _REG_FONT
        ws05.cell(row=req_r, column=3, value=row["wave"]).font = _REG_FONT
        ws05.cell(row=req_r, column=4, value=row["activity"]).font = _REG_FONT
        ws05.cell(row=req_r, column=5, value=row["role"]).font = _BOLD_FONT
        ws05.cell(row=req_r, column=6, value=row["delivery_model"]).font = _REG_FONT
        ws05.cell(row=req_r, column=7, value=row["working_days"]).font = _REG_FONT
        ws05.cell(row=req_r, column=8, value=row["effort_hours"]).font = _REG_FONT
        ws05.cell(row=req_r, column=9, value=row["required_fte"]).font = _REG_FONT
        req_r += 1
    _autofit_columns(ws05)

    # --------------------------------------------------------------------------
    # Sheet 06: Monthly FTE
    # --------------------------------------------------------------------------
    ws06 = wb.create_sheet("06 Monthly FTE")
    ws06["A1"] = "Monthly Required FTE by Role"
    ws06["A1"].font = _TITLE_FONT

    headers06 = ["Functional Role", "Category"] + month_keys + ["Total PD", "Average FTE", "Peak FTE"]
    for c_i, h in enumerate(headers06, start=1):
        ws06.cell(row=3, column=c_i, value=h)
    _style_header_row(ws06, 3, len(headers06))

    m_start_col = 3
    m_end_col = m_start_col + len(month_keys) - 1
    row_06 = 4

    for r_def in STANDARD_ROLES:
        title = r_def.title
        ws06.cell(row=row_06, column=1, value=title).font = _BOLD_FONT
        ws06.cell(row=row_06, column=2, value=r_def.category).font = _REG_FONT

        for m_i, m_key in enumerate(month_keys):
            fte = role_month_fte.get(title, {}).get(m_key, 0.0)
            c = ws06.cell(row=row_06, column=m_start_col + m_i, value=fte)
            c.number_format = "0.00"
            c.font = _REG_FONT

        tot_pd = role_demand.get("by_role_pd", {}).get(title, 0.0)
        c_tot = ws06.cell(row=row_06, column=m_end_col + 1, value=tot_pd)
        c_tot.number_format = "0.0"
        c_tot.font = _BOLD_FONT

        c_avg = ws06.cell(row=row_06, column=m_end_col + 2,
                          value=f"=AVERAGE({get_column_letter(m_start_col)}{row_06}:{get_column_letter(m_end_col)}{row_06})"
                          if month_keys else 0.0)
        c_avg.number_format = "0.00"
        c_avg.font = _REG_FONT

        c_peak = ws06.cell(row=row_06, column=m_end_col + 3,
                           value=f"=MAX({get_column_letter(m_start_col)}{row_06}:{get_column_letter(m_end_col)}{row_06})"
                           if month_keys else 0.0)
        c_peak.number_format = "0.00"
        c_peak.font = _BOLD_FONT

        row_06 += 1

    # Total row across roles
    ws06.cell(row=row_06, column=1, value="Total Team Required FTE").font = _BOLD_FONT
    for m_i in range(len(month_keys)):
        c_col = m_start_col + m_i
        col_let = get_column_letter(c_col)
        c_sum = ws06.cell(row=row_06, column=c_col, value=f"=SUM({col_let}4:{col_let}{row_06-1})")
        c_sum.number_format = "0.00"
        c_sum.font = _BOLD_FONT

    ws06.cell(row=row_06, column=m_end_col + 1,
              value=f"=SUM({get_column_letter(m_end_col + 1)}4:{get_column_letter(m_end_col + 1)}{row_06-1})").font = _BOLD_FONT
    _autofit_columns(ws06)

    # --------------------------------------------------------------------------
    # Sheet 07: Capacity Heatmap
    # --------------------------------------------------------------------------
    ws07 = wb.create_sheet("07 Capacity Heatmap")
    ws07["A1"] = "Role Capacity & Utilization Heatmap"
    ws07["A1"].font = _TITLE_FONT
    ws07["A2"] = "Numeric utilization % based on supplied capacity. Pre-sales defaults to unassessed."
    ws07["A2"].font = _ITALIC_FONT

    headers07 = ["Role Title", "Capacity Mode"] + month_keys
    for c_i, h in enumerate(headers07, start=1):
        ws07.cell(row=4, column=c_i, value=h)
    _style_header_row(ws07, 4, len(headers07))

    h_map = {item["role"]: item for item in cap.get("capacity_heatmap", [])}
    row_07 = 5
    for r_def in STANDARD_ROLES:
        title = r_def.title
        ws07.cell(row=row_07, column=1, value=title).font = _BOLD_FONT
        ws07.cell(row=row_07, column=2, value="Supplied" if cap.get("has_capacity_input") else "Not Provided").font = _REG_FONT

        item = h_map.get(title)
        m_cells = {c["month"]: c for c in item.get("monthly_utilization", [])} if item else {}

        for m_i, m_key in enumerate(month_keys):
            cell_info = m_cells.get(m_key, {})
            u_pct = cell_info.get("utilization_pct")
            st = cell_info.get("status", "UNASSESSED")

            c = ws07.cell(row=row_07, column=3 + m_i)
            if u_pct is not None:
                c.value = u_pct / 100.0
                c.number_format = "0.0%"
                if st == "OVER_ALLOCATED":
                    c.fill = PatternFill("solid", fgColor=_DANGER_BG)
                elif st == "OPTIMAL":
                    c.fill = PatternFill("solid", fgColor=_SUCCESS_BG)
                else:
                    c.fill = PatternFill("solid", fgColor=_LIGHT_BG)
            else:
                c.value = "Unassessed"
                c.font = _ITALIC_FONT
                c.fill = PatternFill("solid", fgColor=_LIGHT_BG)

        row_07 += 1
    _autofit_columns(ws07)

    # --------------------------------------------------------------------------
    # Sheet 08: Wave Loading
    # --------------------------------------------------------------------------
    ws08 = wb.create_sheet("08 Wave Loading")
    ws08["A1"] = "Resource Demand by Migration Wave"
    ws08["A1"].font = _TITLE_FONT

    headers08 = ["Wave Name", "Wave Kind", "Servers Migrated", "Applications", "Estimated Adopt PD", "Share of Factory %", "Key Technical Roles"]
    for c_i, h in enumerate(headers08, start=1):
        ws08.cell(row=3, column=c_i, value=h)
    _style_header_row(ws08, 3, len(headers08))

    row_08 = 4
    for w in demand.get("wave_loading", []):
        ws08.cell(row=row_08, column=1, value=w["wave"]).font = _BOLD_FONT
        ws08.cell(row=row_08, column=2, value=w["kind"]).font = _REG_FONT
        ws08.cell(row=row_08, column=3, value=w["servers"]).font = _REG_FONT
        ws08.cell(row=row_08, column=4, value=w["apps"]).font = _REG_FONT
        ws08.cell(row=row_08, column=5, value=w["pd"]).font = _REG_FONT
        c_sh = ws08.cell(row=row_08, column=6, value=w["share_pct"] / 100.0)
        c_sh.number_format = "0.0%"
        ws08.cell(row=row_08, column=7, value=", ".join(w["critical_roles"])).font = _ITALIC_FONT
        row_08 += 1

    if demand.get("wave_loading"):
        ws08.cell(row=row_08, column=1, value="Total").font = _BOLD_FONT
        ws08.cell(row=row_08, column=3, value=f"=SUM(C4:C{row_08-1})").font = _BOLD_FONT
        ws08.cell(row=row_08, column=4, value=f"=SUM(D4:D{row_08-1})").font = _BOLD_FONT
        ws08.cell(row=row_08, column=5, value=f"=SUM(E4:E{row_08-1})").font = _BOLD_FONT
    _autofit_columns(ws08)

    # --------------------------------------------------------------------------
    # Sheet 09: Role x Phase Matrix
    # --------------------------------------------------------------------------
    ws09 = wb.create_sheet("09 Role x Phase Matrix")
    ws09["A1"] = "Role Effort Across MEG Lifecycle Phases (Person-Days)"
    ws09["A1"].font = _TITLE_FONT

    phases = ["Strategy", "Plan", "Ready", "Adopt", "Govern", "Manage"]
    headers09 = ["Role Title"] + phases + ["Total PD"]
    for c_i, h in enumerate(headers09, start=1):
        ws09.cell(row=3, column=c_i, value=h)
    _style_header_row(ws09, 3, len(headers09))

    role_phase = demand.get("role_phase_matrix", {})
    row_09 = 4
    for r_def in STANDARD_ROLES:
        title = r_def.title
        ws09.cell(row=row_09, column=1, value=title).font = _BOLD_FONT
        for p_i, p in enumerate(phases, start=2):
            val = role_phase.get(title, {}).get(p, 0.0)
            ws09.cell(row=row_09, column=p_i, value=val).font = _REG_FONT

        p_end = get_column_letter(1 + len(phases))
        ws09.cell(row=row_09, column=2 + len(phases), value=f"=SUM(B{row_09}:{p_end}{row_09})").font = _BOLD_FONT
        row_09 += 1

    # Total row across roles
    ws09.cell(row=row_09, column=1, value="Phase Total PD").font = _BOLD_FONT
    for p_i in range(2, 2 + len(phases) + 1):
        c_let = get_column_letter(p_i)
        ws09.cell(row=row_09, column=p_i, value=f"=SUM({c_let}4:{c_let}{row_09-1})").font = _BOLD_FONT
    _autofit_columns(ws09)

    # --------------------------------------------------------------------------
    # Sheet 10: Resource Cost
    # --------------------------------------------------------------------------
    ws10 = wb.create_sheet("10 Resource Cost")
    ws10["A1"] = "Commercial Services Cost by Role & Delivery Model"
    ws10["A1"].font = _TITLE_FONT

    headers10 = ["Functional Role", "Delivery Model", "Billable PD", "Day Rate", "Currency", "Total Cost"]
    for c_i, h in enumerate(headers10, start=1):
        ws10.cell(row=3, column=c_i, value=h)
    _style_header_row(ws10, 3, len(headers10))

    r_costs = {rc["role"]: rc for rc in comm.get("role_costs", [])}
    row_10 = 4
    for r_def in STANDARD_ROLES:
        title = r_def.title
        rc = r_costs.get(title, {})
        pd = rc.get("person_days", 0.0)
        rate = rc.get("day_rate", 0.0)

        ws10.cell(row=row_10, column=1, value=title).font = _BOLD_FONT
        ws10.cell(row=row_10, column=2, value=rc.get("delivery_model", "Onshore")).font = _REG_FONT
        ws10.cell(row=row_10, column=3, value=pd).font = _REG_FONT
        c_rate = ws10.cell(row=row_10, column=4, value=rate)
        c_rate.number_format = "#,##0"
        ws10.cell(row=row_10, column=5, value=comm.get("currency", "USD")).font = _REG_FONT

        c_cost = ws10.cell(row=row_10, column=6, value=f"=C{row_10}*D{row_10}")
        c_cost.number_format = "#,##0"
        c_cost.font = _BOLD_FONT
        row_10 += 1

    ws10.cell(row=row_10, column=1, value="Total Commercial Services").font = _BOLD_FONT
    ws10.cell(row=row_10, column=3, value=f"=SUM(C4:C{row_10-1})").font = _BOLD_FONT
    c_tot_cost = ws10.cell(row=row_10, column=6, value=f"=SUM(F4:F{row_10-1})")
    c_tot_cost.number_format = "#,##0"
    c_tot_cost.font = _BOLD_FONT
    _autofit_columns(ws10)

    # --------------------------------------------------------------------------
    # Sheet 11: DACI-RACI
    # --------------------------------------------------------------------------
    ws11 = wb.create_sheet("11 DACI-RACI")
    ws11["A1"] = "Migration Governance Matrix (RACI & DACI Baseline)"
    ws11["A1"].font = _TITLE_FONT
    ws11["A2"] = "DRAFT — CUSTOMER VALIDATION REQUIRED. Functional roles only, no personal names."
    ws11["A2"].font = _ITALIC_FONT

    ws11["A4"] = "MEG Lifecycle Phase"
    ws11["B4"] = "Accountable (A)"
    ws11["C4"] = "Responsible (R)"
    ws11["D4"] = "Consulted (C)"
    ws11["E4"] = "Informed (I)"
    _style_header_row(ws11, 4, 5)

    raci_baseline = [
        ("Strategy", "Executive Sponsor", "Migration Programme Manager", "Lead Cloud Architect, Security Lead", "Application Owners"),
        ("Plan & Discovery", "Migration Programme Manager", "Lead Cloud Architect, Migration Engineer", "App Owner, DBA Lead", "Executive Sponsor"),
        ("Ready & Landing Zone", "Lead Cloud Architect", "Network Engineer, Security Lead", "DevOps Lead", "Migration Programme Manager"),
        ("Adopt & Migration Waves", "Migration Programme Manager", "Migration Engineer, DBA Lead", "Lead Cloud Architect, App Owner", "Executive Sponsor"),
        ("Govern & Manage", "Lead Cloud Architect", "DevOps Lead, Security Lead", "FinOps Analyst", "Migration Programme Manager"),
    ]
    for r_i, row in enumerate(raci_baseline, start=5):
        for c_i, v in enumerate(row, start=1):
            ws11.cell(row=r_i, column=c_i, value=v).font = _REG_FONT
    _autofit_columns(ws11)

    # --------------------------------------------------------------------------
    # Sheet 12: Skill Gaps
    # --------------------------------------------------------------------------
    ws12 = wb.create_sheet("12 Skill Gaps")
    ws12["A1"] = "Skill Gap & Capacity Risk Analysis"
    ws12["A1"].font = _TITLE_FONT

    headers12 = ["Role Title", "Critical Skill Required", "Required Peak FTE", "Supplied Capacity FTE", "Deficit (FTE)", "When Needed", "Affected Waves", "Mitigation Recommendation"]
    for c_i, h in enumerate(headers12, start=1):
        ws12.cell(row=3, column=c_i, value=h)
    _style_header_row(ws12, 3, len(headers12))

    gaps = cap.get("skill_gaps", [])
    row_12 = 4
    if gaps:
        for g in gaps:
            ws12.cell(row=row_12, column=1, value=g["role"]).font = _BOLD_FONT
            ws12.cell(row=row_12, column=2, value=g["required_skill"]).font = _REG_FONT
            ws12.cell(row=row_12, column=3, value=g["required_fte"]).font = _REG_FONT
            ws12.cell(row=row_12, column=4, value=g["available_fte"]).font = _REG_FONT
            ws12.cell(row=row_12, column=5, value=g["gap_fte"]).font = _BOLD_FONT
            ws12.cell(row=row_12, column=6, value=g["when_needed"]).font = _REG_FONT
            ws12.cell(row=row_12, column=7, value=", ".join(g["affected_waves"])).font = _REG_FONT
            ws12.cell(row=row_12, column=8, value=g["recommendation"]).font = _ITALIC_FONT
            row_12 += 1
    else:
        ws12.cell(row=4, column=1, value="No skill gaps detected or capacity not supplied.").font = _ITALIC_FONT
    _autofit_columns(ws12)

    # --------------------------------------------------------------------------
    # Sheet 13: Calculation Appendix
    # --------------------------------------------------------------------------
    ws13 = wb.create_sheet("13 Calculation Appendix")
    ws13["A1"] = "Resource Demand Calculation Appendix & Traceability"
    ws13["A1"].font = _TITLE_FONT

    headers13 = ["Ref", "Metric / Output", "Result Value", "Unit", "Mathematical Formula / Derivation", "Source Inputs", "Confidence"]
    for c_i, h in enumerate(headers13, start=1):
        ws13.cell(row=3, column=c_i, value=h)
    _style_header_row(ws13, 3, len(headers13))

    app_data = [
        ("R1", "Delivery Person-Days Subtotal", totals.get("delivery_subtotal_pd"), "PD", "Sum of mobilisation, assessment, LZ, execution, testing, cutover, and hypercare", f"Servers={params.get('server_count')}, Apps={params.get('app_count')}", params.get("dq_confidence")),
        ("R2", "Programme Management", totals.get("pm_pd"), "PD", "Delivery Subtotal × 15%", "R1 × 0.15", params.get("dq_confidence")),
        ("R3", "Governance & Assurance", totals.get("governance_pd"), "PD", "Delivery Subtotal × 10%", "R1 × 0.10", params.get("dq_confidence")),
        ("R4", "Contingency Buffer", totals.get("contingency_pd"), "PD", "Base Delivery × Contingency % (from DQ score)", "DQ Confidence", params.get("dq_confidence")),
        ("R5", "Estimate at Completion (EAC)", totals.get("total_person_days"), "PD", "R1 + R2 + R3 + R4", "All workstream PDs", params.get("dq_confidence")),
        ("R6", "Peak Team Size", totals.get("peak_fte"), "FTE", f"Max(Monthly Required FTE) = Max(Monthly PD ÷ {params.get('working_days_per_month', 21)})", "Wave Schedule", params.get("dq_confidence")),
    ]
    for r_i, row in enumerate(app_data, start=4):
        for c_i, val in enumerate(row, start=1):
            ws13.cell(row=r_i, column=c_i, value=val).font = _REG_FONT
    _autofit_columns(ws13)

    # --------------------------------------------------------------------------
    # Sheet 14: Demand vs Capacity
    # --------------------------------------------------------------------------
    ws14 = wb.create_sheet("14 Demand vs Capacity")
    ws14["A1"] = "Role-by-Role Demand vs Supplied Capacity"
    ws14["A1"].font = _TITLE_FONT

    headers14 = ["Functional Role", "Category", "Required Avg FTE", "Required Peak FTE", "Available FTE", "Variance (FTE)", "Peak Utilization", "Period Needed"]
    for c_i, h in enumerate(headers14, start=1):
        ws14.cell(row=3, column=c_i, value=h)
    _style_header_row(ws14, 3, len(headers14))

    d_vs_c = cap.get("demand_vs_capacity", [])
    row_14 = 4
    for item in d_vs_c:
        ws14.cell(row=row_14, column=1, value=item["role"]).font = _BOLD_FONT
        ws14.cell(row=row_14, column=2, value=item["category"]).font = _REG_FONT
        ws14.cell(row=row_14, column=3, value=item["required_avg_fte"]).font = _REG_FONT
        ws14.cell(row=row_14, column=4, value=item["required_peak_fte"]).font = _REG_FONT
        ws14.cell(row=row_14, column=5, value=item["available_display"]).font = _REG_FONT

        c_var = ws14.cell(row=row_14, column=6, value=item["gap_display"])
        c_var.font = _BOLD_FONT
        if isinstance(item["gap_display"], (int, float)) and item["gap_display"] < 0:
            c_var.fill = PatternFill("solid", fgColor=_DANGER_BG)

        ws14.cell(row=row_14, column=7, value=f"{item['utilization_display']}%" if isinstance(item["utilization_display"], (int, float)) else item["utilization_display"]).font = _REG_FONT
        ws14.cell(row=row_14, column=8, value=item["when_needed"]).font = _ITALIC_FONT
        row_14 += 1
    _autofit_columns(ws14)

    # --------------------------------------------------------------------------
    # Sheet 15: Assumption Register
    # --------------------------------------------------------------------------
    ws15 = wb.create_sheet("15 Assumption Register")
    ws15["A1"] = "Migration Resource Planning Assumptions & Risk Register"
    ws15["A1"].font = _TITLE_FONT

    headers15 = ["Ref ID", "Category", "Statement / Assumption Description", "Source", "Impact / Consequence"]
    for c_i, h in enumerate(headers15, start=1):
        ws15.cell(row=3, column=c_i, value=h)
    _style_header_row(ws15, 3, len(headers15))

    reg_items = [
        ("A1", "Assumption", "Customer provides dedicated Application Owners/SMEs throughout wave UAT windows.", "SOW", "Defects addressed within wave soak"),
        ("A2", "Assumption", "Source infrastructure credentials and hypervisor read access provided within 5 business days.", "SOW", "Replication setup without delay"),
        ("A3", "Assumption", "Network routing and ExpressRoute hybrid connectivity established prior to Wave 1 kickoff.", "Ready", "Landing zone cutover readiness"),
        ("X1", "Exclusion", "Refactoring legacy application code or rewriting database stored procedures.", "Scope", "Legacy refactoring requires separate project"),
        ("X2", "Exclusion", "Physical hardware decommissioning or rack dismantling.", "Scope", "Customer facility management responsibility"),
        ("G1", "Data Gap", f"Estate assessment completed with {params.get('dq_confidence')} data quality confidence.", "Inventory", f"Contingency buffered at {round(totals.get('contingency_pd', 0))} PD"),
    ]
    for r_i, row in enumerate(reg_items, start=4):
        for c_i, val in enumerate(row, start=1):
            ws15.cell(row=r_i, column=c_i, value=val).font = _REG_FONT
    _autofit_columns(ws15)

    # Save to in-memory bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
