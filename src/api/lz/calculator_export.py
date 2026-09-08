"""
Parse an Azure Pricing Calculator Excel export (`ExportedEstimate.xlsx`) back into
structured data (PRD E11.16 / E11.17).

    parsed = parse_calculator_export(path_or_bytes)
    # -> {estimate_name, currency, licensing_program, line_items[], support_monthly,
    #     total_monthly, total_upfront, created_at, source}

The calculator's export is one sheet (named after the estimate), fixed shape:

    row 1   Microsoft Azure Estimate
    row 2   <estimate name>
    row 3   Service category | Service type | Custom name | Region | Description |
            Estimated monthly cost | Estimated upfront cost
    row 4…  one row per configured resource
            Support        | … | 0
            (blank) | | | Licensing Program | Microsoft Customer Agreement (MCA)
            (blank) | | | Billing Account  |
            (blank) | | | Billing Profile  |
            (blank) | | | Total | | <monthly> | <upfront>
            Disclaimer
            All prices shown are in United States Dollar ($) …
            This estimate was created at 9/8/2026 12:40:29 PM UTC …

Pure. openpyxl only.
"""
from __future__ import annotations

import io
import re

_HEADER_A = "service category"
_STOP_D = {"support", "licensing program", "billing account", "billing profile",
           "total", "disclaimer"}


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return round(float(str(v).replace(",", "").replace("$", "").strip()), 2)
    except (TypeError, ValueError):
        return None


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def parse_calculator_export(source) -> dict:
    """`source` is a path, a file-like, or raw `bytes`."""
    from openpyxl import load_workbook

    buf = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
    wb = load_workbook(buf, read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    def cell(r, c):
        return rows[r][c] if r < len(rows) and c < len(rows[r]) else None

    estimate_name = _s(cell(1, 0)) or _s(ws.title)

    # locate the header row
    hdr = next((i for i, r in enumerate(rows)
                if _s(r[0]).lower() == _HEADER_A), 2)

    line_items: list[dict] = []
    support_monthly = 0.0
    licensing_program = ""
    total_monthly = total_upfront = None
    created_at = ""
    currency = ""

    for r in rows[hdr + 1:]:
        a, b, cname, region, desc, mo, up = (list(r) + [None] * 7)[:7]
        da, dd = _s(a).lower(), _s(region).lower()

        if dd == "support" or da == "support":
            support_monthly = _num(mo) or 0.0
            continue
        if dd == "licensing program":
            licensing_program = _s(desc)
            continue
        if dd == "total":
            total_monthly, total_upfront = _num(mo), _num(up)
            continue
        if da == "disclaimer" or dd in ("billing account", "billing profile"):
            continue
        if da.startswith("all prices shown are in"):
            m = re.search(r"in\s+(.*?)\s*(?:\(([^)]+)\))?\s*([A-Z]{3})?\s*\.?$", _s(a))
            if m:
                currency = (m.group(3) or "").strip() or _s(a)
            mcur = re.search(r"\b([A-Z]{3})\b", _s(a))
            if mcur:
                currency = mcur.group(1)
            continue
        if da.startswith("this estimate was created at"):
            created_at = _s(a).split("created at", 1)[-1].strip()
            continue
        if not any([_s(a), _s(b), _s(desc)]):
            continue

        line_items.append({
            "category": _s(a),
            "service_type": _s(b),
            "custom_name": _s(cname),
            "region": _s(region),
            "description": _s(desc),
            "monthly": _num(mo) or 0.0,
            "upfront": _num(up) or 0.0,
        })

    if total_monthly is None:
        total_monthly = round(sum(li["monthly"] for li in line_items) + support_monthly, 2)
    if total_upfront is None:
        total_upfront = round(sum(li["upfront"] for li in line_items), 2)

    return {
        "estimate_name": estimate_name,
        "currency": currency or "USD",
        "licensing_program": licensing_program,
        "line_items": line_items,
        "line_count": len(line_items),
        "support_monthly": support_monthly,
        "total_monthly": total_monthly,
        "total_upfront": total_upfront,
        "annual": round((total_monthly or 0) * 12, 2),
        "created_at": created_at,
        "source": "azure-pricing-calculator-export",
    }


def reconcile(parsed: dict, internal_monthly: float | None) -> dict:
    """Compare the calculator POE total against Landfall's own run-rate."""
    poe = parsed.get("total_monthly")
    if not internal_monthly or not poe:
        return {"delta_pct": None, "note": "one side missing — no comparison"}
    delta = (poe - internal_monthly) / internal_monthly * 100
    return {
        "internal_monthly": round(internal_monthly, 2),
        "poe_monthly": round(poe, 2),
        "delta_pct": round(delta, 1),
        "within_tolerance": abs(delta) <= 15,
        "note": ("calculator POE and Landfall's internal estimate agree within 15%"
                 if abs(delta) <= 15 else
                 "calculator POE and Landfall's internal estimate differ by more than "
                 "15% — reconcile the line items before submitting"),
    }
