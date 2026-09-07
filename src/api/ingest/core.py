"""
Pure ingestion logic — no Azure imports, unit-testable on its own.

    headers, rows = read_table("servers.csv", data_bytes)
    result = normalize("servers.csv", data_bytes)   # -> NormResult

A NormResult carries the target table, the normalised rows (dicts whose keys are
`scripts/schema.sql` columns), the profile that matched, the source headers that
were never mapped, and any row-level issues found while coercing values.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Callable, NamedTuple, Optional

# ---------------------------------------------------------------------------
# reading .csv / .xlsx into (headers, list[dict])
# ---------------------------------------------------------------------------
def read_table(name: str, data: bytes) -> tuple[list[str], list[dict]]:
    lower = name.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    return _read_csv(data)


def _read_csv(data: bytes) -> tuple[list[str], list[dict]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip() for h in (reader.fieldnames or [])]
    rows = []
    for raw in reader:
        rows.append({(k.strip() if k else k): v for k, v in raw.items()})
    return headers, rows


def _read_xlsx(data: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook  # available in the Functions runtime

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    names = wb.sheetnames
    preferred = [n for n in names if re.search(r"vinfo|vm_?info|server|inventory|host", n, re.I)]
    ws = wb[preferred[0]] if preferred else wb[names[0]]

    it = ws.iter_rows(values_only=True)
    try:
        first = next(it)
    except StopIteration:
        return [], []
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(first)]
    rows = []
    for raw in it:
        if raw is None or all(c is None for c in raw):
            continue
        rows.append({headers[i]: (raw[i] if i < len(raw) else None) for i in range(len(headers))})
    return headers, rows


# ---------------------------------------------------------------------------
# value transforms
# ---------------------------------------------------------------------------
def _txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _num(v):
    if v is None or v == "":
        return None
    m = re.match(r"^\s*-?\d[\d,]*(\.\d+)?", str(v))
    return float(m.group().replace(",", "").strip()) if m else None


def _int(v):
    n = _num(v)
    return int(round(n)) if n is not None else None


def _mib_to_gib(v):
    n = _num(v)
    return round(n / 1024, 2) if n is not None else None


def _mb_to_gb(v):
    n = _num(v)
    return round(n / 1000, 2) if n is not None else None


def _bit(v):
    s = (_txt(v) or "").lower()
    if s in ("1", "true", "yes", "y", "t"):
        return 1
    if s in ("0", "false", "no", "n", "f", ""):
        return 0 if s else None
    return None


def _powerstate(v):
    s = (_txt(v) or "").lower()
    if not s:
        return None
    if any(k in s for k in ("off", "stopped", "deallocat", "halt")):
        return "poweredOff"
    if any(k in s for k in ("on", "running", "started", "up")):
        return "poweredOn"
    return _txt(v)


def _iso_date(v):
    s = _txt(v)
    if not s:
        return None
    s = s.split("T")[0].split(" ")[0]
    for pat, fmt in (
        (r"^\d{4}-\d{2}-\d{2}$", "%Y-%m-%d"),
        (r"^\d{2}/\d{2}/\d{4}$", "%m/%d/%Y"),
        (r"^\d{1,2}-[A-Za-z]{3}-\d{2,4}$", None),
    ):
        if re.match(pat, s):
            if fmt:
                from datetime import datetime

                try:
                    return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
                except ValueError:
                    return None
            return s
    return s if re.match(r"^\d{4}-\d{2}-\d{2}$", s) else None


_OS_FAMILIES = [
    (r"windows", "Windows Server"),
    (r"red\s?hat|rhel", "Red Hat Enterprise Linux"),
    (r"ubuntu", "Ubuntu"),
    (r"suse|sles", "SUSE Linux Enterprise Server"),
    (r"cent\s?os", "CentOS"),
    (r"oracle\s+linux", "Oracle Linux"),
    (r"debian", "Debian"),
    (r"amazon\s+linux", "Amazon Linux"),
]


def _os_name(v):
    s = _txt(v)
    if not s:
        return None
    for pat, fam in _OS_FAMILIES:
        if re.search(pat, s, re.I):
            return fam
    return s.split("(")[0].strip()[:128]


def _os_version(v):
    s = _txt(v)
    if not s:
        return None
    m = re.search(r"(20\d{2}(?:\s?R\d)?|\d{1,2}(?:\.\d{1,2}){0,2}(?:\s?(?:LTS|SP\d))?)", s)
    return m.group(1).strip()[:64] if m else None


# ---------------------------------------------------------------------------
# source profiles
# ---------------------------------------------------------------------------
class Col(NamedTuple):
    sources: tuple[str, ...]            # candidate headers, lower-case; "~x" = substring match
    transform: Callable = _txt
    required: bool = False


class Profile(NamedTuple):
    name: str
    table: str                          # target table in scripts/schema.sql
    signature: tuple[tuple[str, ...], ...]   # groups; a group matches if any header hits
    columns: dict


def _norm_keys(row: dict) -> dict:
    return {(str(k).strip().lower() if k is not None else ""): v for k, v in row.items()}


def _pick(nrow: dict, candidates: tuple[str, ...]):
    """Return (value, matched_header_lower) for the first candidate present."""
    for c in candidates:
        if c.startswith("~"):
            sub = c[1:]
            for k, v in nrow.items():
                if sub in k:
                    return v, k
        elif c in nrow:
            return nrow[c], c
    return None, None


# --- Landfall-native CSVs (columns already match schema.sql) ---
_NATIVE_SERVERS = {
    c: Col((c,), t)
    for c, t in {
        "server_id": _txt, "hostname": _txt, "env": _txt, "os_name": _txt,
        "os_version": _txt, "os_eol_date": _iso_date, "vcpu": _int, "ram_gb": _num,
        "provisioned_disk_gb": _num, "used_disk_gb": _num, "cpu_avg_pct": _num,
        "cpu_peak_pct": _num, "ram_avg_pct": _num, "disk_iops_avg": _num,
        "disk_iops_peak": _num, "net_in_gb_30d": _num, "net_out_gb_30d": _num,
        "cluster": _txt, "datacenter": _txt,
        "powerstate": _powerstate, "app_id": _txt, "notes": _txt,
    }.items()
}
_NATIVE_PERF = {
    "server_id": Col(("server_id", "host", "vm"), _txt, required=True),
    "sample_date": Col(("sample_date", "date", "timestamp", "day", "collected"), _iso_date, required=True),
    "cpu_avg_pct": Col(("cpu_avg_pct", "cpu avg", "cpu_avg"), _num),
    "cpu_peak_pct": Col(("cpu_peak_pct", "cpu peak", "cpu_max"), _num),
    "cpu_p95_pct": Col(("cpu_p95_pct", "cpu p95"), _num),
    "mem_avg_pct": Col(("mem_avg_pct", "memory avg", "mem_avg", "ram_avg_pct"), _num),
    "mem_peak_pct": Col(("mem_peak_pct", "memory peak", "mem_max"), _num),
    "mem_p95_pct": Col(("mem_p95_pct", "mem p95"), _num),
    "disk_iops_avg": Col(("disk_iops_avg", "iops avg", "iops"), _num),
    "disk_iops_peak": Col(("disk_iops_peak", "iops peak", "iops_max"), _num),
    "disk_read_iops_avg": Col(("disk_read_iops_avg", "read iops"), _num),
    "disk_write_iops_avg": Col(("disk_write_iops_avg", "write iops"), _num),
    "disk_throughput_mbps_avg": Col(("disk_throughput_mbps_avg", "disk mbps", "throughput mbps"), _num),
    "net_in_gb": Col(("net_in_gb", "network in gb", "ingress gb", "rx gb"), _num),
    "net_out_gb": Col(("net_out_gb", "network out gb", "egress gb", "tx gb"), _num),
    "net_in_peak_mbps": Col(("net_in_peak_mbps", "net in peak", "rx peak mbps"), _num),
    "net_out_peak_mbps": Col(("net_out_peak_mbps", "net out peak", "tx peak mbps"), _num),
}
_NATIVE_APPS = {
    c: Col((c,), t)
    for c, t in {
        "app_id": _txt, "app_name": _txt, "business_owner": _txt, "criticality": _int,
        "users": _int, "tech_stack": _txt, "db_engine": _txt, "internet_facing": _bit,
        "compliance_scope": _txt, "disposition": _txt, "complexity": _txt, "wave": _int,
    }.items()
}
_NATIVE_DEPS = {
    c: Col((c,), t)
    for c, t in {
        "src_id": _txt, "dst_id": _txt, "port": _int, "protocol": _txt,
        "direction": _txt, "confidence": _txt, "bytes_30d_gb": _num,
        "flows_30d": _int, "last_seen": _iso_date,
    }.items()
}
_NATIVE_STORAGE = {
    c: Col((c,), t)
    for c, t in {
        "storage_id": _txt, "server_id": _txt, "type": _txt, "size_gb": _num,
        "iops": _int, "target_service": _txt,
    }.items()
}

# --- RVTools vInfo tab ---
_RVTOOLS = {
    "hostname": Col(("vm", "~vm name"), _txt, required=True),
    "powerstate": Col(("powerstate", "power state", "power_state"), _powerstate),
    "vcpu": Col(("cpus", "cpu", "~cpu count"), _int),
    "ram_gb": Col(("memory", "~memory mib", "~memory"), _mib_to_gib),
    "provisioned_disk_gb": Col(("provisioned mib", "provisioned mb", "provisioned", "~provisioned"), _mib_to_gib),
    "used_disk_gb": Col(("in use mib", "in use mb", "~in use"), _mib_to_gib),
    "os_name": Col(("os according to the configuration file", "os according to the vmware tools", "~os according"), _os_name),
    "os_version": Col(("os according to the configuration file", "os according to the vmware tools", "~os according"), _os_version),
    "cluster": Col(("cluster",), _txt),
    "datacenter": Col(("datacenter", "data center"), _txt),
    "notes": Col(("annotation", "notes"), _txt),
}

# --- generic CMDB (ServiceNow / BMC style) export ---
_CMDB = {
    "hostname": Col(("name", "fqdn", "host name", "hostname", "ci name"), _txt, required=True),
    "env": Col(("environment", "used for", "u_environment", "stage"), lambda v: (_txt(v) or "").lower() or None),
    "vcpu": Col(("cpu count", "cpus", "cpu", "num cpu", "processor count", "cpu_count"), _int),
    "ram_gb": Col(("ram", "memory", "ram (gb)", "memory (gb)"), _num),
    "provisioned_disk_gb": Col(("disk space", "disk", "storage", "disk (gb)", "disk_space"), _num),
    "os_name": Col(("os", "operating system", "os name"), _os_name),
    "os_version": Col(("os version", "os_version", "version"), _os_version),
    "datacenter": Col(("data center", "datacenter", "location", "u_datacenter", "site"), _txt),
    "powerstate": Col(("status", "operational status", "state", "power"), _powerstate),
    "notes": Col(("short description", "comments", "description"), _txt),
}


PROFILES: tuple[Profile, ...] = (
    Profile("landfall_performance", "performance",
            (("server_id", "host"), ("sample_date", "date", "timestamp"),
             ("cpu_avg_pct", "cpu avg")), _NATIVE_PERF),
    Profile("landfall_servers", "servers",
            (("server_id", "hostname"), ("vcpu",), ("ram_gb",)), _NATIVE_SERVERS),
    Profile("landfall_applications", "applications",
            (("app_id",), ("app_name",)), _NATIVE_APPS),
    Profile("landfall_dependencies", "dependencies",
            (("src_id",), ("dst_id",)), _NATIVE_DEPS),
    Profile("landfall_storage", "storage",
            (("storage_id",), ("server_id",), ("size_gb",)), _NATIVE_STORAGE),
    Profile("rvtools_vinfo", "servers",
            (("vm",), ("cpus",), ("memory", "~memory"), ("powerstate", "power state")), _RVTOOLS),
    Profile("generic_cmdb", "servers",
            (("name", "fqdn", "ci name"), ("cpu count", "cpus", "cpu", "cpu_count"),
             ("ram", "memory")), _CMDB),
)

_FILENAME_HINTS = {
    "perf": "landfall_performance", "utilization": "landfall_performance",
    "utilisation": "landfall_performance", "metrics": "landfall_performance",
    "server": "landfall_servers", "vminfo": "rvtools_vinfo", "vinfo": "rvtools_vinfo",
    "rvtools": "rvtools_vinfo", "application": "landfall_applications",
    "app": "landfall_applications", "portfolio": "landfall_applications",
    "dependenc": "landfall_dependencies", "netflow": "landfall_dependencies",
    "storage": "landfall_storage", "volume": "landfall_storage", "disk": "landfall_storage",
    "cmdb": "generic_cmdb",
}
_BY_NAME = {p.name: p for p in PROFILES}


def _signature_score(profile: Profile, header_set: set[str]) -> float:
    hits = 0
    for group in profile.signature:
        if any((g[1:] in h if g.startswith("~") else g == h) for g in group for h in header_set):
            hits += 1
    return hits / len(profile.signature)


def detect_profile(name: str, headers: list[str]) -> Optional[Profile]:
    header_set = {h.strip().lower() for h in headers if h}
    lname = name.lower()

    # 1. filename hint, if its signature at least half-matches
    for token, pname in _FILENAME_HINTS.items():
        if token in lname:
            p = _BY_NAME[pname]
            if _signature_score(p, header_set) >= 0.5:
                return p

    # 2. best signature match, native profiles winning ties
    scored = sorted(
        ((p, _signature_score(p, header_set)) for p in PROFILES),
        key=lambda ps: (ps[1], ps[0].name.startswith("landfall")),
        reverse=True,
    )
    best, score = scored[0]
    return best if score >= 0.6 else None


# ---------------------------------------------------------------------------
# normalize
# ---------------------------------------------------------------------------
class Issue(NamedTuple):
    level: str          # "error" | "warning"
    where: str
    message: str


class NormResult(NamedTuple):
    table: Optional[str]
    rows: list[dict]
    profile: Optional[str]
    unmapped_headers: list[str]
    issues: list[Issue]
    row_count_in: int


def normalize(name: str, data: bytes) -> NormResult:
    headers, raw_rows = read_table(name, data)
    profile = detect_profile(name, headers)
    if profile is None:
        return NormResult(
            None, [], None, headers,
            [Issue("error", name, "unrecognised file — no source profile matched its headers")],
            len(raw_rows),
        )

    unmapped = {h.strip() for h in headers if h}
    issues: list[Issue] = []
    out: list[dict] = []

    for i, row in enumerate(raw_rows):
        nrow = _norm_keys(row)
        rec: dict = {}
        for target, col in profile.columns.items():
            value, matched = _pick(nrow, col.sources)
            if matched:
                # remove any original-case header that lower-cases to `matched`
                for h in list(unmapped):
                    if h.lower() == matched:
                        unmapped.discard(h)
            try:
                rec[target] = col.transform(value)
            except Exception as exc:  # noqa: BLE001 - never let one bad cell kill the file
                rec[target] = None
                issues.append(Issue("warning", f"{profile.table} row {i + 2}",
                                    f"could not parse {target}={value!r} ({exc})"))
            if col.required and rec.get(target) in (None, ""):
                issues.append(Issue("error", f"{profile.table} row {i + 2}",
                                    f"missing required column {target}"))
        rec["source_file"] = name
        out.append(rec)

    # synthesise a stable id where the schema needs one but the source has none
    _fill_ids(profile.table, out, name)

    return NormResult(profile.table, out, profile.name, sorted(unmapped), issues, len(raw_rows))


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (s or "").lower()).strip("-")[:40] or "row"


def _fill_ids(table: str, rows: list[dict], name: str) -> None:
    base = _slug(name.rsplit(".", 1)[0])
    if table == "servers":
        for n, r in enumerate(rows, 1):
            if not r.get("server_id"):
                r["server_id"] = r.get("hostname") or f"{base}-srv-{n:04d}"
    elif table == "applications":
        for n, r in enumerate(rows, 1):
            if not r.get("app_id"):
                r["app_id"] = _slug(r.get("app_name") or f"{base}-app-{n:04d}")
    elif table == "storage":
        for n, r in enumerate(rows, 1):
            if not r.get("storage_id"):
                r["storage_id"] = f"{base}-stg-{n:05d}"
