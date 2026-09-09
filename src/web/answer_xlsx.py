"""Turn a chat answer (+ any tabular tool output it carried) into an Excel
workbook — PRD E11.14 "Ask & export to Excel".

    build_answer_workbook(engagement, question, answer, tables=[...], sql="...")

Sheets: `Answer` (the Q&A + a DRAFT note), one sheet per table (`columns` + `rows`),
and `Provenance` (engagement, the SQL that produced each table, when it was
generated, and that it is a Landfall DRAFT). Pure — no Azure, no network.
"""
from __future__ import annotations

import datetime as _dt
import io
import re

_MAX_SHEETS = 12
_MAX_ROWS = 5000
_INVALID = re.compile(r"[:\\/?*\[\]]")


def _sheet_name(raw: str, used: set[str]) -> str:
    name = _INVALID.sub("-", str(raw or "table")).strip() or "table"
    name = name[:31]
    base, i = name, 2
    while name.lower() in used:
        suffix = f" ({i})"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def build_answer_workbook(engagement: str | None, question: str, answer: str,
                          tables: list[dict] | None = None, sql: str | None = None,
                          generated_at: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    ts = generated_at or _dt.datetime.now(_dt.timezone.utc).isoformat(timespec="seconds")
    tables = tables or []

    wb = Workbook()
    ans = wb.active
    ans.title = "Answer"
    bold = Font(bold=True)
    ans["A1"] = "Landfall — inventory answer (DRAFT)"
    ans["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Engagement", engagement or "—"),
        ("Generated (UTC)", ts),
        ("Question", question or "—"),
        ("Answer", answer or "—"),
        ("Note", "DRAFT for architect review — figures carry the assumptions in the "
                 "source estimate. Not a bid."),
    ]
    r = 3
    for k, v in rows:
        ans.cell(row=r, column=1, value=k).font = bold
        c = ans.cell(row=r, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ans.column_dimensions["A"].width = 18
    ans.column_dimensions["B"].width = 110

    used = {"answer"}
    prov_rows: list[tuple] = []
    for t in tables[:_MAX_SHEETS - 2]:
        cols = list(t.get("columns") or [])
        data = list(t.get("rows") or [])
        title = _sheet_name(t.get("name") or "results", used)
        ws = wb.create_sheet(title)
        if cols:
            ws.append([str(c) for c in cols])
            for cell in ws[1]:
                cell.font = bold
        for row in data[:_MAX_ROWS]:
            ws.append(["" if v is None else v for v in row])
        if len(data) > _MAX_ROWS:
            ws.append([f"… {len(data) - _MAX_ROWS} more rows not exported"])
        prov_rows.append((title, len(data), t.get("sql") or sql or ""))

    prov = wb.create_sheet("Provenance")
    prov["A1"] = "Provenance"
    prov["A1"].font = Font(bold=True, size=13)
    prov.append([])
    for k, v in (("Engagement", engagement or "—"), ("Generated (UTC)", ts),
                 ("Source", "Landfall migration estimator — query_inventory over Azure SQL"),
                 ("Status", "DRAFT")):
        prov.append([k, v])
        prov.cell(row=prov.max_row, column=1).font = bold
    if prov_rows:
        prov.append([])
        prov.append(["Sheet", "Rows", "SQL"])
        for cell in prov[prov.max_row]:
            cell.font = bold
        for name, n, q in prov_rows:
            prov.append([name, n, q])
    elif sql:
        prov.append([])
        prov.append(["SQL", sql])
        prov.cell(row=prov.max_row, column=1).font = bold
    prov.column_dimensions["A"].width = 20
    prov.column_dimensions["B"].width = 90
    prov.column_dimensions["C"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tables_from_response(resp) -> tuple[list[dict], str | None]:
    """Best-effort: pull query_inventory results out of a Responses API result so the
    answer can be exported to Excel. Returns (tables, last_sql). Tolerant of shape
    drift — an unrecognised response just yields ([], None)."""
    import json as _json

    tables: list[dict] = []
    last_sql = None
    names: dict[str, str] = {}
    for item in getattr(resp, "output", None) or []:
        itype = getattr(item, "type", "") or ""
        if itype in ("function_call", "tool_call"):
            cid = getattr(item, "call_id", None) or getattr(item, "id", None)
            nm = getattr(item, "name", None) or ""
            if cid:
                names[cid] = nm
        if itype in ("function_call_output", "tool_call_output", "function_call_result"):
            cid = getattr(item, "call_id", None) or getattr(item, "id", None)
            nm = names.get(cid, "")
            out = getattr(item, "output", None)
            if out is None:
                continue
            try:
                payload = _json.loads(out) if isinstance(out, str) else out
            except Exception:  # noqa: BLE001
                continue
            if not isinstance(payload, dict):
                continue
            cols, rows = payload.get("columns"), payload.get("rows")
            if isinstance(cols, list) and isinstance(rows, list):
                sql = payload.get("sql")
                if sql:
                    last_sql = sql
                if nm in ("", "query_inventory") or "quer" in nm.lower():
                    tables.append({"name": payload.get("name") or "results",
                                   "columns": cols, "rows": rows, "sql": sql})
    return tables, last_sql
