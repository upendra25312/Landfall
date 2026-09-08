"""
Upload validation + a best-effort peek for the engagement Upload panel
(PRD E11.6 / E11.24). Pure — no Azure, no network.

  ok, kind, reason = classify(filename, first_bytes)
  info = peek(filename, data)      # {rows, columns, profile}  (best effort)

Types are checked by **magic bytes + structure**, not by trusting the extension.
The importer in src/api/ingest does the real profiling at analysis time; this is
only the instant "✓ uploaded — RVTools vInfo · 412 rows" confirmation.
"""
from __future__ import annotations

import csv
import io

MAX_FILE = 100 * 1024 * 1024          # 100 MB per file
MAX_REQUEST = 250 * 1024 * 1024       # 250 MB per upload request
MAX_ENGAGEMENT = 2 * 1024 * 1024 * 1024  # 2 GB per engagement (soft — warn)

# extension -> (destination folder, accepted magic prefixes | None for text)
_DATA = "inventory"
_DOCS = "docs"
_EXT: dict[str, tuple[str, tuple[bytes, ...] | None]] = {
    ".csv": (_DATA, None),
    ".tsv": (_DATA, None),
    ".json": (_DATA, None),
    ".xlsx": (_DATA, (b"PK\x03\x04",)),
    ".xls": (_DATA, (b"\xd0\xcf\x11\xe0",)),
    ".zip": (_DATA, (b"PK\x03\x04", b"PK\x05\x06")),
    ".pdf": (_DOCS, (b"%PDF",)),
    ".docx": (_DOCS, (b"PK\x03\x04",)),
    ".md": (_DOCS, None),
    ".txt": (_DOCS, None),
    ".png": (_DOCS, (b"\x89PNG\r\n\x1a\n",)),
    ".jpg": (_DOCS, (b"\xff\xd8\xff",)),
    ".jpeg": (_DOCS, (b"\xff\xd8\xff",)),
}
_BLOCKED = {".xlsm", ".xltm", ".docm", ".dotm", ".pptm", ".exe", ".dll", ".bat",
            ".cmd", ".sh", ".ps1", ".js", ".vbs", ".jar", ".msi", ".7z", ".rar",
            ".gz", ".tar"}

ACCEPT_ATTR = ",".join(sorted(_EXT))   # for the <input accept="…">


def _ext(name: str) -> str:
    name = (name or "").strip().lower()
    return name[name.rfind("."):] if "." in name else ""


def safe_name(name: str) -> str:
    """basename only, keep [A-Za-z0-9._-], collapse the rest."""
    import re
    base = (name or "").replace("\\", "/").split("/")[-1]
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._") or "file"
    return base[:120]


def classify(filename: str, head: bytes) -> tuple[bool, str, str]:
    """(accepted, kind, reason). kind is 'inventory' | 'docs' when accepted."""
    ext = _ext(filename)
    if not ext:
        return False, "", "no file extension - rename it with .csv / .xlsx / .pdf ..."
    if ext in _BLOCKED:
        if ext in (".xlsm", ".xltm", ".docm", ".dotm", ".pptm"):
            return False, "", f"macro-enabled Office files aren't accepted - re-save as {ext[:-1]}x"
        return False, "", f"{ext} files aren't accepted"
    spec = _EXT.get(ext)
    if not spec:
        return False, "", f"{ext} isn't a supported type"
    kind, magics = spec
    if magics is not None:
        if not head or not any(head.startswith(m) for m in magics):
            return False, "", f"the file doesn't look like a real {ext} (content check failed)"
    else:
        # text formats — must decode and not be binary
        sample = head[:4096]
        if b"\x00" in sample:
            return False, "", f"{ext} should be text but looks binary"
        try:
            sample.decode("utf-8")
        except UnicodeDecodeError:
            try:
                sample.decode("latin-1")
            except UnicodeDecodeError:
                return False, "", f"{ext} isn't readable as text"
    return True, kind, ""


# --- best-effort peek -------------------------------------------------------

_HINTS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("RVTools vInfo", ("vm", "powerstate", "cpus")),
    ("RVTools vInfo", ("vm", "power state", "in use mib")),
    ("server inventory", ("hostname", "vcpu", "ram_gb")),
    ("server inventory", ("server_id", "vcpu")),
    ("CMDB / CI export", ("ci name", "cpu count", "environment")),
    ("CMDB / CI export", ("name", "cpu", "memory", "environment")),
    ("application portfolio", ("app_id", "app_name")),
    ("application portfolio", ("application", "criticality")),
    ("dependencies / flows", ("src_id", "dst_id")),
    ("dependencies / flows", ("source", "destination", "port")),
    ("storage / volumes", ("storage_id", "size_gb")),
    ("storage / volumes", ("volume", "capacity")),
    ("performance / utilisation", ("server_id", "cpu_avg_pct")),
    ("performance / utilisation", ("host", "date", "cpu")),
)


def _hint(headers: list[str]) -> str:
    hset = {h.strip().lower() for h in headers}
    best, score = "", 0.0
    for label, sig in _HINTS:
        hits = sum(1 for s in sig if s in hset)
        if hits and hits / len(sig) > score:
            best, score = label, hits / len(sig)
    return best if score >= 0.5 else ""


def peek(filename: str, data: bytes) -> dict:
    """{profile, columns, rows} — best effort, never raises."""
    ext = _ext(filename)
    try:
        if ext in (".csv", ".tsv"):
            text = data.decode("utf-8", "replace")
            dialect_delim = "\t" if ext == ".tsv" else None
            sample = text[:65536]
            if dialect_delim is None:
                dialect_delim = "\t" if sample.count("\t") > sample.count(",") else ","
            rdr = csv.reader(io.StringIO(text), delimiter=dialect_delim)
            headers = next(rdr, []) or []
            rows = sum(1 for _ in rdr)
            return {"profile": _hint(headers), "columns": len(headers), "rows": rows}
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            headers = [str(c) for c in (next(it, ()) or ()) if c is not None]
            rows = sum(1 for _ in it)
            wb.close()
            return {"profile": _hint(headers), "columns": len(headers), "rows": rows}
        if ext == ".json":
            import json as _j
            obj = _j.loads(data.decode("utf-8", "replace"))
            if isinstance(obj, list):
                return {"profile": "", "columns": len(obj[0]) if obj and isinstance(obj[0], dict) else 0,
                        "rows": len(obj)}
            return {"profile": "", "columns": len(obj) if isinstance(obj, dict) else 0, "rows": 0}
    except Exception:  # noqa: BLE001
        pass
    return {"profile": "", "columns": 0, "rows": 0}
