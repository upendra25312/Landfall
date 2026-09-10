"""C52: chat, persisted history, and Excel export cross the extracted routers."""
import os
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect

from .test_chat_page import _goto, _select


def test_chat_reload_and_excel_export(page, base_url, console_errors, tmp_path):
    _goto(page, base_url)
    _select(page, "contoso-ltd/dc-exit")
    question = "Describe this engagement for the router regression check"
    page.fill("#q", question)
    page.click("#send")
    expect(page.locator("#log")).to_contain_text("Confidence: Medium")
    page.reload(wait_until="networkidle")
    expect(page.locator("#log")).to_contain_text(question)
    expect(page.locator("#log")).to_contain_text("Confidence: Medium")
    with page.expect_download() as pending:
        page.get_by_role("button", name="Download as Excel").last.click()
    download = pending.value
    target = tmp_path / download.suggested_filename
    download.save_as(target)
    workbook = load_workbook(target)
    assert "Answer" in workbook.sheetnames and "Provenance" in workbook.sheetnames
    assert any(question in str(cell.value) for row in workbook["Answer"] for cell in row)
    workbook.close()
    assert console_errors == []
    # Optional retained before/after evidence; normal CI needs no output directory.
    if folder := os.environ.get("C52_SCREENSHOT_DIR"):
        destination = Path(folder)
        destination.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(destination / "chat-desktop.png"), full_page=True)
        page.set_viewport_size({"width": 375, "height": 812})
        page.screenshot(path=str(destination / "chat-mobile.png"), full_page=True)
