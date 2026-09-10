"""C53: browser journeys for create/upload/analysis and specific downloads."""
import io
import os
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect

from .test_chat_page import _goto, _select


def test_create_upload_analyse_and_prompt_card(page, base_url, console_errors):
    _goto(page, base_url)
    page.click("#neweng")
    page.fill("input[name=customer]", "Validation C53")
    page.fill("input[name=project]", "Browser journey")
    page.click("#ef button[type=submit]")
    expect(page.locator("#engsel")).to_have_value("validation-c53/browser-journey")
    expect(page.locator("#pipeline .step.done")).to_have_count(0)
    page.locator("#ufile").set_input_files({"name": "servers.csv", "mimeType": "text/csv", "buffer": (
        b"server_id,hostname,env,os_name,vcpu,ram_gb,provisioned_disk_gb,used_disk_gb,powerstate\n"
        b"s1,web01,prod,Ubuntu,2,16,128,90,poweredOn\n")})
    expect(page.locator("#filerows")).to_contain_text("servers.csv")
    expect(page.locator("#startanalysis")).to_be_visible()
    with page.expect_response(lambda r: r.url.endswith("/analyze") and r.request.method == "POST") as response:
        page.click("#startanalysis")
    assert response.value.status == 200
    expect(page.locator("#dqsummary")).to_contain_text("1 rows")
    expect(page.locator("#pipeline .step.done")).to_have_count(2)
    page.get_by_role("button", name="Full estimate").click()
    expect(page.locator("#log")).to_contain_text("Here is the full estimate")
    expect(page.get_by_role("button", name="Download as Excel")).to_be_visible()
    page.reload(wait_until="networkidle")
    expect(page.locator("#log")).to_contain_text("Here is the full estimate")
    assert console_errors == []
    if folder := os.environ.get("C53_SCREENSHOT_DIR"):
        Path(folder).mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(Path(folder) / "create-upload-analysis.png"), full_page=True)


def test_empty_engagement_does_not_display_another_estate_as_published(page, base_url, console_errors):
    _goto(page, base_url)
    _select(page, "northwind/pilot")
    expect(page.locator("#pipeline .step.done")).to_have_count(0)
    response = page.request.get(base_url + "/dashboard/data?e=northwind/pilot")
    assert response.status == 404
    questionnaire = page.request.get(base_url + "/questionnaire.xlsx?e=northwind/pilot")
    assert questionnaire.status == 200
    workbook = load_workbook(io.BytesIO(questionnaire.body()))
    assert workbook.sheetnames
    workbook.close()
    assert console_errors == []
