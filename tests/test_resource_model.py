"""
Unit and Integration Tests for Deterministic Resource Demand & Capacity Planning (PRD E15.3 / E15C).

Verifies:
  - Role catalogue integrity and MEG alignment
  - Determinism of demand derivation
  - Exact mathematical reconciliation with deliverable.effort.estimate_effort
  - Calendar support (dated vs. relative months)
  - Zero hallucination when capacity is unsupplied (Available = NOT PROVIDED, Gap = UNKNOWN)
  - Capacity variance and RESOURCE-CONSTRAINED SCHEDULE warnings
  - Skill-gap analysis accuracy
  - Conservative, Expected, and Accelerated scenario differences
  - Commercial cost modeling and rate overrides
  - SOW resource narrative and 12-column project plan generation
  - 15-sheet openpyxl workbook structure, formulas, and recalculation compatibility
"""
from __future__ import annotations

import io
import os
import sys

import pytest
from conftest import ROOT

sys.path.insert(0, os.path.join(ROOT, "src", "api"))

from openpyxl import load_workbook
from resource.capacity import analyze_capacity_and_constraints
from resource.commercial import calculate_commercial_cost
from resource.engine import derive_resource_demand
from resource.model import ROLES_BY_ID, ROLES_BY_TITLE, STANDARD_ROLES
from resource.scenarios import generate_resource_scenarios
from resource.sow_project_plan import build_project_plan, build_sow_resource_section
from resource.workbook import generate_resource_workbook
from waves.schedule import build_schedule


@pytest.fixture(scope="module")
def sample_inputs():
    servers = 250
    apps = 31
    disps = [{"disposition": "Rehost"}] * 25 + [{"disposition": "Replatform"}] * 3 + [{"disposition": "Repurchase"}] * 2 + [{"disposition": "Retire"}] * 1
    wp = {"waves": [
        {"wave": 1, "kind": "platform", "server_count": 4, "app_count": 1},
        {"wave": 2, "kind": "pilot", "server_count": 16, "app_count": 3},
        {"wave": 3, "kind": "app", "server_count": 38, "app_count": 5},
        {"wave": 4, "kind": "app", "server_count": 42, "app_count": 6},
        {"wave": 5, "kind": "app", "server_count": 48, "app_count": 6},
        {"wave": 6, "kind": "app", "server_count": 52, "app_count": 5},
        {"wave": 7, "kind": "regulated", "server_count": 50, "app_count": 5},
    ]}
    sched = build_schedule(wp, start_date="2026-10-05")
    return {
        "server_count": servers,
        "app_count": apps,
        "dispositions": disps,
        "wave_plan": wp,
        "schedule": sched,
        "spokes": 9,
        "regulated": True,
        "dq_confidence": "Medium",
    }


def test_standard_roles_catalogue():
    assert len(STANDARD_ROLES) >= 10
    role_ids = {r.id for r in STANDARD_ROLES}
    assert "prog_manager" in role_ids
    assert "lead_architect" in role_ids
    assert "migration_engineer" in role_ids
    assert "dba_lead" in role_ids
    assert "security_lead" in role_ids
    assert "network_engineer" in role_ids
    assert "app_owner" in role_ids
    assert "test_lead" in role_ids
    assert "devops_ops_lead" in role_ids
    assert "finops_analyst" in role_ids

    # Each role must have category, skills, and default delivery model
    for r in STANDARD_ROLES:
        assert r.title in ROLES_BY_TITLE
        assert r.id in ROLES_BY_ID
        assert len(r.skills) >= 2
        assert r.category


def test_derive_resource_demand_deterministic(sample_inputs):
    d1 = derive_resource_demand(**sample_inputs)
    d2 = derive_resource_demand(**sample_inputs)
    assert d1["totals"] == d2["totals"]
    assert d1["role_demand"]["by_role_pd"] == d2["role_demand"]["by_role_pd"]
    assert d1["reconciliation"] == d2["reconciliation"]
    assert len(d1["requirements"]) == len(d2["requirements"])


def test_demand_reconciles_to_effort_engine(sample_inputs):
    d = derive_resource_demand(**sample_inputs)
    recon = d["reconciliation"]
    assert recon["reconciled"] is True
    assert abs(recon["variance"]) < 0.01

    # Check totals
    totals = d["totals"]
    assert totals["total_person_days"] > 500.0
    assert totals["total_hours"] == round(totals["total_person_days"] * 8.0, 1)
    assert totals["peak_fte"] > 0
    assert totals["avg_fte"] > 0
    assert totals["peak_fte"] >= totals["avg_fte"]


def test_calendar_support_relative_and_dated(sample_inputs):
    # 1. Dated schedule
    d_dated = derive_resource_demand(**sample_inputs)
    m_dated = d_dated["monthly_schedule"]
    assert m_dated["has_dated_schedule"] is True
    assert any("-" in k for k in m_dated["month_keys"])  # YYYY-MM format

    # 2. Undated / Relative calendar
    inputs_undated = dict(sample_inputs)
    inputs_undated["schedule"] = None
    d_undated = derive_resource_demand(**inputs_undated)
    m_undated = d_undated["monthly_schedule"]
    assert m_undated["has_dated_schedule"] is False
    assert all(k.startswith("Month ") for k in m_undated["month_keys"])


def test_unsupplied_capacity_reports_not_provided_zero_hallucination(sample_inputs):
    d = derive_resource_demand(**sample_inputs)
    cap = analyze_capacity_and_constraints(d, capacity_input=None)

    assert cap["has_capacity_input"] is False
    assert cap["status"] == "CAPACITY NOT PROVIDED"
    assert cap["is_constrained"] is False
    assert cap["constraints_count"] == 0
    assert len(cap["skill_gaps"]) == 0

    # All rows must report NOT PROVIDED / UNKNOWN, never invented numbers
    for row in cap["demand_vs_capacity"]:
        assert row["available_fte"] is None
        assert row["available_display"] == "NOT PROVIDED"
        assert row["variance"] is None
        assert row["gap_display"] == "UNKNOWN"
        assert row["utilization_pct"] is None
        assert row["utilization_display"] == "UNKNOWN"


def test_supplied_capacity_satisfies_or_flags_constraints(sample_inputs):
    d = derive_resource_demand(**sample_inputs)

    # 1. Constrained capacity scenario (only 1 engineer available, but peak requires > 3)
    constrained_input = {
        "Infrastructure / Migration Engineer": 1.0,
        "Lead Cloud Architect": 1.0,
        "Migration Programme Manager": 1.0,
    }
    cap_c = analyze_capacity_and_constraints(d, capacity_input=constrained_input)
    assert cap_c["has_capacity_input"] is True
    assert cap_c["is_constrained"] is True
    assert cap_c["status"] == "RESOURCE-CONSTRAINED SCHEDULE"
    assert cap_c["constraints_count"] > 0
    assert len(cap_c["bottleneck_months"]) > 0

    eng_row = next(r for r in cap_c["demand_vs_capacity"] if r["role"] == "Infrastructure / Migration Engineer")
    assert eng_row["available_fte"] == 1.0
    assert eng_row["variance"] < 0  # deficit
    assert eng_row["utilization_pct"] > 100.0

    # 2. Generous capacity scenario (10 engineers available)
    satisfied_input = {
        "Infrastructure / Migration Engineer": 10.0,
        "Lead Cloud Architect": 5.0,
        "Migration Programme Manager": 2.0,
        "Database Administrator (DBA)": 5.0,
        "Security & Compliance Lead": 3.0,
        "Network Engineer": 3.0,
        "Test Lead": 3.0,
        "Application Owner / SME": 5.0,
        "DevOps & Operations Lead": 2.0,
        "FinOps Analyst": 2.0,
        "Change Manager": 2.0,
    }
    cap_s = analyze_capacity_and_constraints(d, capacity_input=satisfied_input)
    assert cap_s["has_capacity_input"] is True
    assert cap_s["is_constrained"] is False
    assert cap_s["status"] == "CAPACITY SATISFIED"
    assert cap_s["constraints_count"] == 0


def test_skill_gap_analysis(sample_inputs):
    d = derive_resource_demand(**sample_inputs)
    constrained_input = {
        "Infrastructure / Migration Engineer": 0.5,
        "Database Administrator (DBA)": 0.2,
    }
    cap = analyze_capacity_and_constraints(d, capacity_input=constrained_input)
    gaps = cap["skill_gaps"]
    assert len(gaps) >= 2
    roles_with_gaps = {g["role"] for g in gaps}
    assert "Infrastructure / Migration Engineer" in roles_with_gaps
    assert "Database Administrator (DBA)" in roles_with_gaps

    for g in gaps:
        assert g["gap_fte"] > 0
        assert g["required_skill"]
        assert "procure" in g["recommendation"].lower() or "partner" in g["recommendation"].lower()


def test_three_scenarios_conservative_expected_accelerated(sample_inputs):
    res = generate_resource_scenarios(**sample_inputs)
    scenarios = res["scenarios"]

    cons = scenarios["conservative"]
    exp = scenarios["expected"]
    acc = scenarios["accelerated"]

    # Throughput ordering: Conservative < Expected < Accelerated
    assert cons["assumptions"]["throughput_servers_per_week"] < exp["assumptions"]["throughput_servers_per_week"]
    assert exp["assumptions"]["throughput_servers_per_week"] < acc["assumptions"]["throughput_servers_per_week"]

    # Contingency ordering: Accelerated (8%) < Expected (12%) < Conservative (20%)
    assert acc["assumptions"]["contingency_pct"] <= exp["assumptions"]["contingency_pct"]
    assert exp["assumptions"]["contingency_pct"] <= cons["assumptions"]["contingency_pct"]

    # Parallel lanes: Accelerated has >= 2 lanes
    assert acc["assumptions"]["parallel_waves"] >= 2
    assert cons["assumptions"]["parallel_waves"] == 1

    # Duration: Accelerated finishes in fewer or equal months compared to Conservative
    assert acc["duration_months"] <= cons["duration_months"]

    # Peak FTE: Accelerated packs work in parallel, so peak FTE is highest
    assert acc["peak_fte"] >= exp["peak_fte"]


def test_commercial_cost_derivation(sample_inputs):
    d = derive_resource_demand(**sample_inputs)
    comm = calculate_commercial_cost(d)

    assert comm["total_commercial_cost"] > 0
    assert comm["currency"] == "USD"
    assert comm["billable_person_days"] > 0
    assert comm["non_billable_person_days"] > 0  # Client SME & Change Manager are $0

    # Custom override test
    overrides = {"Lead Cloud Architect": 1500.0, "Infrastructure / Migration Engineer": 900.0}
    comm_ov = calculate_commercial_cost(d, rates_override=overrides)
    assert comm_ov["day_rates"]["Lead Cloud Architect"] == 1500.0
    assert comm_ov["day_rates"]["Infrastructure / Migration Engineer"] == 900.0


def test_sow_and_project_plan_generation(sample_inputs):
    d = derive_resource_demand(**sample_inputs)
    comm = calculate_commercial_cost(d)
    sow = build_sow_resource_section(d, comm)

    assert len(sow["roles"]) >= 10
    assert len(sow["staffing_assumptions"]) >= 3
    assert len(sow["customer_responsibilities"]) >= 3
    assert len(sow["partner_responsibilities"]) >= 3
    assert len(sow["exclusions"]) >= 3

    plan = build_project_plan(d, sample_inputs["schedule"])
    assert len(plan) >= 12
    # Check 12 columns exist on tasks
    first = plan[0]
    expected_cols = {"id", "phase", "workstream", "task", "start", "finish", "duration_weeks", "dependency", "owner_role", "wave", "status", "milestone"}
    assert expected_cols.issubset(set(first.keys()))


def test_openpyxl_workbook_generation(sample_inputs):
    d = derive_resource_demand(**sample_inputs)
    cap = analyze_capacity_and_constraints(d, capacity_input={"Lead Cloud Architect": 2.0})
    comm = calculate_commercial_cost(d)
    sow = build_sow_resource_section(d, comm)
    plan = build_project_plan(d, sample_inputs["schedule"])

    wb_bytes = generate_resource_workbook(
        demand=d,
        capacity_analysis=cap,
        commercial=comm,
        sow_narrative=sow,
        project_plan=plan,
    )
    assert len(wb_bytes) > 5000

    # Load and inspect sheets
    wb = load_workbook(io.BytesIO(wb_bytes), data_only=False)
    sheet_names = wb.sheetnames
    assert len(sheet_names) == 15

    expected_sheets = [
        "01 Executive Summary",
        "02 Resource Assumptions",
        "03 Role Catalogue",
        "04 Activities",
        "05 Resource Demand",
        "06 Monthly FTE",
        "07 Capacity Heatmap",
        "08 Wave Loading",
        "09 Role x Phase Matrix",
        "10 Resource Cost",
        "11 DACI-RACI",
        "12 Skill Gaps",
        "13 Calculation Appendix",
        "14 Demand vs Capacity",
        "15 Assumption Register",
    ]
    for s in expected_sheets:
        assert s in sheet_names, f"Missing sheet: {s}"

    assert wb.calculation.fullCalcOnLoad is True
    assert wb.calculation.forceFullCalc is True

    # Verify formula syntax in sheet 06 and 10
    ws06 = wb["06 Monthly FTE"]
    formulas_06 = [c.value for row in ws06.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas_06) > 0
    for f in formulas_06:
        assert f.startswith("=SUM(") or f.startswith("=AVERAGE(") or f.startswith("=MAX(")

    ws10 = wb["10 Resource Cost"]
    formulas_10 = [c.value for row in ws10.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("*D" in f for f in formulas_10)
    assert any("SUM(" in f for f in formulas_10)
