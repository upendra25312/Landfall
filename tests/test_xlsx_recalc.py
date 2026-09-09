"""C22 / E11.12 (E5.4q) — the estimate workbook is a *live model*: derived cells
are formulas that reflow, and they recalc without error.

The pure checks always run. The LibreOffice recalc runs only when RECALC=1 and
`soffice` is on PATH (the CI job sets both); it opens the workbook, forces a
recalculation, and asserts no formula produced an error.
"""
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile

import pytest
from conftest import ROOT

sys.path.insert(0, os.path.join(ROOT, "src", "api"))
sys.path.insert(0, os.path.join(ROOT, "evals"))

import pipeline as P  # noqa: E402
from deliverable.export import to_xlsx  # noqa: E402
from deliverable.xlsx_model import model_formulas  # noqa: E402

_ERR = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
_FORMULA_OK = re.compile(r"^=[A-Z0-9.+\-*/() ]+$")


@pytest.fixture(scope="module")
def workbook_bytes():
    return to_xlsx(P.run())


def _model_rows(wb_bytes):
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(wb_bytes))["Model"]
    out = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row and row[0]:
            out[row[0]] = row[2]      # figure_id -> Result cell (formula str or number)
    return out


def test_headline_chain_is_formulas_not_literals(workbook_bytes):
    rows = _model_rows(workbook_bytes)
    # the run-rate roll-up and the annualisation must be live
    assert isinstance(rows.get("F8"), str) and rows["F8"].startswith("=")
    assert isinstance(rows.get("F9"), str) and rows["F9"].startswith("=")
    # F9 (annual) references F8 (monthly) — the classic "change monthly, annual follows"
    assert "C12" in rows["F9"] or re.search(r"C\d+\*12", rows["F9"])


def test_every_formula_is_excel_2007_safe(workbook_bytes):
    rows = _model_rows(workbook_bytes)
    formulas = [v for v in rows.values() if isinstance(v, str) and v.startswith("=")]
    assert formulas, "no formulas in the model sheet"
    for f in formulas:
        assert _FORMULA_OK.match(f), f"unsafe / non-Excel-2007 formula: {f}"


def test_modelled_figures_reproduce_their_stated_result():
    for m in model_formulas(P.run()):
        assert abs(m["expected"] - m["result"]) <= max(0.02, 0.01 * abs(m["result"])), m


def test_at_least_the_cost_rollup_is_modelled():
    ids = {m["figure_id"] for m in model_formulas(P.run())}
    assert {"F8", "F9"} <= ids            # run-rate monthly + annual are always derivable


# ---------------------------------------------------------------- LibreOffice
_soffice = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(os.environ.get("RECALC") != "1" or not _soffice,
                    reason="set RECALC=1 and install LibreOffice to run the recalc gate")
def test_libreoffice_recalc_has_no_errors(workbook_bytes):
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "estimate.xlsx")
        with open(src, "wb") as fh:
            fh.write(workbook_bytes)
        subprocess.run([_soffice, "--headless", "--calc", "--convert-to",
                        "xlsx:Calc MS Excel 2007 XML", "--outdir", d, src],
                       check=True, timeout=180, capture_output=True)
        out = os.path.join(d, "estimate.xlsx")   # LO overwrites in place name
        from openpyxl import load_workbook
        wb = load_workbook(out, data_only=True)
        bad = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and v in _ERR:
                        bad.append((ws.title, v))
        assert not bad, f"recalc produced errors: {bad}"

        # the recalculated run-rate matches what the estimator computed
        model = load_workbook(out, data_only=True)["Model"]
        vals = {r[0]: r[2] for r in model.iter_rows(min_row=5, values_only=True) if r and r[0]}
        want = {m["figure_id"]: m["result"] for m in model_formulas(P.run())}
        for fid in ("F8", "F9"):
            assert isinstance(vals.get(fid), (int, float)), (fid, vals.get(fid))
            assert abs(vals[fid] - want[fid]) <= max(0.5, 0.01 * want[fid]), (fid, vals[fid], want[fid])
