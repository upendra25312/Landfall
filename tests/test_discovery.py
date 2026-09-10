"""C20b / E11.25 — discovery questionnaire: served, exportable to Word/Excel,
re-imported into _discovery.json, and folded into the estimate's assumptions.
"""
import datetime as _dt
import importlib
import io
import json
import os
import sys

import pytest
from conftest import ROOT

sys.path.insert(0, os.path.join(ROOT, "src", "web"))
sys.path.insert(0, os.path.join(ROOT, "src", "api"))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import discovery as D  # noqa: E402


# --------------------------------------------------------------- catalog / drift

def test_catalog_is_in_sync_with_the_html():
    gen = importlib.import_module("gen_discovery_catalog")
    assert gen.main(check=True) == 0, "run scripts/gen_discovery_catalog.py — catalog drifted"


def test_catalog_shape():
    cat = D.load_catalog()
    assert len(cat) >= 60
    ids = [q["id"] for q in cat]
    assert len(ids) == len(set(ids))                      # unique
    assert {"SC1", "R1", "B4", "N2"} <= set(ids)          # the ones that feed the model
    for q in cat:
        assert q["priority"] in ("MUST", "SHOULD", "NICE")
        assert q["section"] and q["text"]
    assert any(q["feeds"] == "compliance_scope" for q in cat)


# --------------------------------------------------------------- export / import

_ANSWERS = {
    "SC1": "PCI-DSS and ISO 27001",
    "R1": "RPO 1h / RTO 4h for tier 1, 24h for tier 3",
    "B4": "West Europe; data must stay in the EU",
    "N2": "ExpressRoute 1 Gbps, circuit already ordered",
}


def test_xlsx_round_trip():
    blob = D.render_xlsx(_ANSWERS)
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(blob))["Discovery"]
    assert ws["A4"].value == "ID"
    parsed = D.parse_upload(blob, "returned.xlsx")
    assert parsed["is_template"] and parsed["format"] == "xlsx"
    assert parsed["answers"] == _ANSWERS


def test_docx_round_trip():
    blob = D.render_docx(_ANSWERS)
    parsed = D.parse_upload(blob, "returned.docx")
    assert parsed["is_template"] and parsed["format"] == "docx"
    assert parsed["answers"] == _ANSWERS


def test_blank_template_exports_do_not_crash():
    assert len(D.render_xlsx(None)) > 2000
    assert len(D.render_docx(None)) > 2000
    # a blank export re-parsed yields no answers and is not a template
    p = D.parse_upload(D.render_xlsx(None), "blank.xlsx")
    assert p["answers"] == {} and not p["is_template"]


def test_non_questionnaire_file_is_not_a_template():
    p = D.parse_upload(b"hostname,vcpu\nweb01,4\n", "servers.csv")
    assert not p["is_template"] and p["answers"] == {}
    p2 = D.parse_upload(b"%PDF-1.4 fake", "scan.pdf")
    assert p2["format"] == "pdf" and not p2["is_template"]


def test_gaps():
    g = D.gaps(_ANSWERS)
    assert g["answered"] == 4
    assert g["must_total"] > g["must_answered"] >= 1
    assert g["headline"].startswith("4/")
    flat = [q["id"] for grp in g["missing"] for q in grp["questions"]]
    assert "SC1" not in flat and "B1" in flat            # answered excluded, unanswered MUST present
    assert all(q["priority"] in ("MUST", "SHOULD") for grp in g["missing"] for q in grp["questions"])


def test_discovery_record_shape():
    rec = D.discovery_record(_ANSWERS, list(_ANSWERS), "q.xlsx", "arch@contoso")
    assert rec["answer_map"] == _ANSWERS
    assert rec["answers"]["SC1"]["feeds"] == "compliance_scope"
    assert rec["answers"]["SC1"]["question"]
    assert rec["imported_by"] == "arch@contoso" and rec["imported_at"]
    assert rec["gaps"]["answered"] == 4


# --------------------------------------------------------------- assemble wiring

def test_assemble_folds_discovery_into_the_register():
    from deliverable.assemble import assemble_estimate
    rec = D.discovery_record(_ANSWERS, list(_ANSWERS), "q.xlsx", "me")
    pkg = assemble_estimate({"inventory_summary": {"servers": 10, "applications": 3},
                             "discovery": rec})
    reg = pkg["register"]
    cited = [a for a in reg["assumptions"] if a["source"].startswith("discovery:")]
    assert {"discovery:SC1", "discovery:R1"} <= {a["source"] for a in cited}
    gap_srcs = {g["source"] for g in reg["data_gaps"]}
    assert any(s.startswith("discovery:") for s in gap_srcs)
    assert reg["discovery"]["answered"] == 4
    # capped at 12 explicit "ask the client" rows + a "+N more" line
    explicit = [g for g in reg["data_gaps"] if g["source"].startswith("discovery:") and g["source"] != "discovery:more"]
    assert len(explicit) <= 12

    from deliverable.export import export
    for fmt in ("docx", "xlsx", "pptx"):
        assert len(export(pkg, fmt)[0]) > 1000


def test_assemble_without_discovery_is_unchanged():
    from deliverable.assemble import assemble_estimate
    pkg = assemble_estimate({"inventory_summary": {"servers": 10, "applications": 3}})
    assert pkg["register"]["discovery"] is None
    assert not [a for a in pkg["register"]["assumptions"] if a["source"].startswith("discovery:")]


# --------------------------------------------------------------- web routes

class _Down:
    def __init__(self, d):
        self._d = d

    def readall(self):
        return self._d


class _Container:
    def __init__(self, store):
        self.store = store

    def get_blob_client(self, key):
        store = self.store

        class _BC:
            def __init__(s):
                s._staged = {}

            def stage_block(s, bid, data):
                s._staged[bid] = bytes(data)

            def commit_block_list(s, blocks, metadata=None, content_settings=None):
                store[key] = {"data": b"".join(s._staged[b.id] for b in blocks),
                              "metadata": dict(metadata or {}),
                              "mtime": _dt.datetime(2026, 9, 9, tzinfo=_dt.timezone.utc)}
        return _BC()

    def download_blob(self, key):
        if key not in self.store:
            raise KeyError(key)
        return _Down(self.store[key]["data"])

    def upload_blob(self, key, data, overwrite=False):
        self.store[key] = {"data": data if isinstance(data, bytes) else bytes(data),
                           "metadata": {}, "mtime": _dt.datetime(2026, 9, 9, tzinfo=_dt.timezone.utc)}

    def list_blobs(self, name_starts_with="", include=None):
        for k, rec in list(self.store.items()):
            if k.startswith(name_starts_with):
                b = type("B", (), {})()
                b.name, b.size, b.last_modified, b.metadata = k, len(rec["data"]), rec["mtime"], rec["metadata"]
                yield b

    def delete_blob(self, key):
        self.store.pop(key, None)


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setenv("FOUNDRY_PROJECT_ENDPOINT", "https://example/api/projects/x")
    monkeypatch.setenv("STORAGE_URL", "https://s.blob.core.windows.net")
    import app as webapp
    import web_storage
    importlib.reload(webapp)
    from fastapi.testclient import TestClient
    store = {
        "engagements/contoso/dc-exit/_engagement.json": {
            "data": json.dumps({"engagement": "contoso/dc-exit", "visibility": "all"}).encode(),
            "metadata": {}, "mtime": _dt.datetime(2026, 9, 9, tzinfo=_dt.timezone.utc)},
    }
    cont = _Container(store)
    monkeypatch.setattr(web_storage, "_raw_container", lambda: cont)
    return webapp, TestClient(webapp.app), store


def test_get_questionnaire_page(client):
    _w, c, _ = client
    r = c.get("/questionnaire")
    assert r.status_code == 200 and "Discovery" in r.text


@pytest.mark.parametrize("fmt", ["xlsx", "docx"])
def test_questionnaire_export_download(client, fmt):
    _w, c, _ = client
    r = c.get(f"/questionnaire.{fmt}")
    assert r.status_code == 200
    assert r.headers["content-disposition"].endswith(f'.{fmt}"')
    assert len(r.content) > 2000


def test_upload_questionnaire_writes_discovery_json(client):
    webapp, c, store = client
    blob = D.render_xlsx(_ANSWERS)
    r = c.post("/api/engagements/contoso/dc-exit/upload",
               files={"file": ("discovery-answers.xlsx", blob,
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
               data={"kind": "docs"})
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["discovery"]["answered"] == 4
    assert "engagements/contoso/dc-exit/_discovery.json" in store
    rec = json.loads(store["engagements/contoso/dc-exit/_discovery.json"]["data"])
    assert rec["answer_map"] == _ANSWERS

    r2 = c.get("/api/engagements/contoso/dc-exit/discovery")
    assert r2.status_code == 200 and r2.json()["imported"] is True
    assert r2.json()["answers"]["SC1"]["answer"] == _ANSWERS["SC1"]


def test_discovery_endpoint_before_any_upload(client):
    _w, c, _ = client
    r = c.get("/api/engagements/contoso/dc-exit/discovery")
    assert r.status_code == 200 and r.json()["imported"] is False
    assert r.json()["gaps"]["answered"] == 0
