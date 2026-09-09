"""C21 / E11.14 — ask & export a chat answer to Excel."""
import io
import json
import os
import sys

from conftest import ROOT

sys.path.insert(0, os.path.join(ROOT, "src", "web"))

from answer_xlsx import build_answer_workbook, tables_from_response  # noqa: E402


def _load(b):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(b))


def test_workbook_has_answer_table_and_provenance_sheets():
    b = build_answer_workbook(
        "contoso/dc-exit", "How many prod Windows servers?",
        "42 servers, 168 vCPU total.",
        [{"name": "prod windows", "columns": ["os", "n"], "rows": [["WS2019", 42]], "sql": "SELECT 1"}],
        sql="SELECT 1")
    wb = _load(b)
    assert wb.sheetnames == ["Answer", "prod windows", "Provenance"]
    ans = wb["Answer"]
    flat = "\n".join(str(c.value) for row in ans.iter_rows() for c in row if c.value)
    assert "contoso/dc-exit" in flat and "DRAFT" in flat and "prod Windows" in flat
    t = wb["prod windows"]
    assert [c.value for c in t[1]] == ["os", "n"]
    assert [c.value for c in t[2]] == ["WS2019", 42]
    prov = "\n".join(str(c.value) for row in wb["Provenance"].iter_rows() for c in row if c.value)
    assert "SELECT 1" in prov and "contoso/dc-exit" in prov


def test_workbook_without_tables_still_valid():
    b = build_answer_workbook("c/p", "q?", "a.", [], None)
    wb = _load(b)
    assert wb.sheetnames == ["Answer", "Provenance"]


def test_sheet_name_is_sanitised_and_deduped():
    b = build_answer_workbook("c/p", "q", "a", [
        {"name": "a/b:c*d?", "columns": ["x"], "rows": [[1]]},
        {"name": "a/b:c*d?", "columns": ["x"], "rows": [[2]]},
    ])
    wb = _load(b)
    assert wb.sheetnames[1] == "a-b-c-d-"
    assert wb.sheetnames[2].startswith("a-b-c-d-") and wb.sheetnames[2] != wb.sheetnames[1]


def test_row_cap_is_enforced():
    rows = [[i] for i in range(5200)]
    b = build_answer_workbook("c/p", "q", "a", [{"name": "big", "columns": ["i"], "rows": rows}])
    ws = _load(b)["big"]
    # 1 header + 5000 data + 1 "more rows" note
    assert ws.max_row == 5002
    assert "not exported" in str(ws.cell(row=5002, column=1).value)


class _Item:
    def __init__(self, **kw):
        self.__dict__.update(kw)


def test_tables_from_response_extracts_query_inventory_output():
    resp = _Item(output=[
        _Item(type="function_call", call_id="c1", name="query_inventory"),
        _Item(type="function_call_output", call_id="c1",
              output=json.dumps({"sql": "SELECT COUNT(*) FROM servers",
                                 "columns": ["n"], "rows": [[250]]})),
        _Item(type="message", content=[]),
    ])
    tables, sql = tables_from_response(resp)
    assert sql == "SELECT COUNT(*) FROM servers"
    assert len(tables) == 1 and tables[0]["rows"] == [[250]]


def test_tables_from_response_tolerates_unknown_shape():
    assert tables_from_response(_Item(output=None)) == ([], None)
    assert tables_from_response(_Item()) == ([], None)
