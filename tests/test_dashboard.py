"""Cycle 14 — assessment dashboard web app + publish_estimate (E5.5)."""
import io
import json
import os
import sys

import pytest
from conftest import ROOT

sys.path.insert(0, os.path.join(ROOT, "src", "web"))
sys.path.insert(0, os.path.join(ROOT, "evals"))

import pipeline as P  # noqa: E402


# --- publish_estimate (Function) --------------------------------------

class _FakeContainer:
    def __init__(self):
        self.blobs = {}

    def upload_blob(self, name, data, overwrite=False):
        self.blobs[name] = data if isinstance(data, bytes) else bytes(data)


def _req(body: dict):
    import azure.functions as func
    return func.HttpRequest(method="POST", url="http://x/api/publish_estimate",
                            body=json.dumps(body).encode(), headers={"Content-Type": "application/json"})


def test_publish_estimate_writes_package_and_three_exports(monkeypatch):
    from deliverable import functions as dfn

    fake = _FakeContainer()
    monkeypatch.setattr(dfn, "_container_client", lambda: fake)

    pkg = P.run()
    resp = dfn.publish_estimate_route(_req({"package": pkg, "engagement": "_default_/_default_"}))
    assert resp.status_code == 200
    out = json.loads(resp.get_body())
    assert set(out["published"]) == {"latest.json", "latest.xlsx", "latest.docx", "latest.pptx"}

    base = "engagements/_default_/_default_/estimate"
    assert set(fake.blobs) == {f"{base}/latest.{x}" for x in ("json", "xlsx", "docx", "pptx")}
    back = json.loads(fake.blobs[f"{base}/latest.json"])
    assert back["meta"]["package_id"] == pkg["meta"]["package_id"]
    assert back["meta"]["engagement"] == "_default_/_default_"
    assert len(back["figures"]) == len(pkg["figures"])

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(fake.blobs[f"{base}/latest.xlsx"]))
    assert "Calculation appendix" in wb.sheetnames


def test_publish_estimate_scopes_by_engagement(monkeypatch):
    from deliverable import functions as dfn
    fake = _FakeContainer()
    monkeypatch.setattr(dfn, "_container_client", lambda: fake)
    pkg = P.run()
    resp = dfn.publish_estimate_route(_req({"package": pkg, "engagement": "contoso-ltd/dc-exit"}))
    assert resp.status_code == 200
    out = json.loads(resp.get_body())
    assert out["engagement"] == "contoso-ltd/dc-exit"
    assert all(n.startswith("engagements/contoso-ltd/dc-exit/estimate/") for n in fake.blobs)
    assert json.loads(fake.blobs["engagements/contoso-ltd/dc-exit/estimate/latest.json"])["meta"]["engagement"] == "contoso-ltd/dc-exit"


class _HistContainer(_FakeContainer):
    """_FakeContainer + read/list so publish_estimate can snapshot the prior version."""
    def download_blob(self, name):
        if name not in self.blobs:
            raise KeyError(name)
        data = self.blobs[name]

        class _D:
            def readall(_s):
                return data
        return _D()

    def list_blobs(self, name_starts_with=""):
        class _B:
            def __init__(_s, n):
                _s.name = n
        return [_B(n) for n in list(self.blobs) if n.startswith(name_starts_with)]


def test_publish_estimate_snapshots_the_prior_version(monkeypatch):
    from deliverable import functions as dfn
    fake = _HistContainer()
    monkeypatch.setattr(dfn, "_container_client", lambda: fake)
    eid = "contoso-ltd/dc-exit"

    r1 = json.loads(dfn.publish_estimate_route(
        _req({"package": P.run(), "engagement": eid})).get_body())
    assert r1["snapshot"] is None                      # nothing to snapshot on the first publish

    r2 = json.loads(dfn.publish_estimate_route(
        _req({"package": P.run(), "engagement": eid})).get_body())
    stamp = r2["snapshot"]
    assert stamp and stamp.endswith("Z")

    hp = f"engagements/{eid}/history/{stamp}"
    assert f"{hp}/latest.json" in fake.blobs
    assert f"{hp}/latest.xlsx" in fake.blobs
    # the live latest.json is the new publish, the snapshot is the old one
    assert json.loads(fake.blobs[f"engagements/{eid}/estimate/latest.json"])["meta"]["published_at"]
    assert json.loads(fake.blobs[f"{hp}/latest.json"])["meta"]["engagement"] == eid


def test_publish_estimate_rejects_empty_body(monkeypatch):
    from deliverable import functions as dfn
    monkeypatch.setattr(dfn, "_container_client", lambda: _FakeContainer())
    resp = dfn.publish_estimate_route(_req({}))
    assert resp.status_code == 400


def test_publish_estimate_build_poe_kicks_calculator_run(monkeypatch):
    """E11.18 hook — build_poe=true stages the POE run after publishing."""
    from deliverable import functions as dfn
    import lz.functions as lzf

    fake = _FakeContainer()
    monkeypatch.setattr(dfn, "_container_client", lambda: fake)
    calls = []
    monkeypatch.setattr(lzf, "stage_calc_run",
                        lambda eng, **kw: calls.append((eng, kw)) or
                        {"engagement": eng, "status": "building", "spec_line_count": 12,
                         "internal_monthly_estimate": 9000.0})

    pkg = P.run()
    body = {"package": pkg, "engagement": "contoso-ltd/dc-exit", "build_poe": True,
            "compute_cost": {"line_items": [{"sku": "x"}]}}   # `raw` non-empty
    resp = dfn.publish_estimate_route(_req(body))
    assert resp.status_code == 200
    out = json.loads(resp.get_body())
    assert out["poe"]["status"] == "building" and out["poe"]["spec_line_count"] == 12
    assert calls == [("contoso-ltd/dc-exit", {"include_dr_compute": False})]


def test_publish_estimate_build_poe_failure_does_not_fail_publish(monkeypatch):
    from deliverable import functions as dfn
    import lz.functions as lzf

    fake = _FakeContainer()
    monkeypatch.setattr(dfn, "_container_client", lambda: fake)

    def _boom(eng, **kw):
        raise lzf._CalcSpecError("no queue wired", 503)

    monkeypatch.setattr(lzf, "stage_calc_run", _boom)
    resp = dfn.publish_estimate_route(
        _req({"package": P.run(), "engagement": "_default_/_default_",
              "build_poe": True, "storage_cost": {"line_items": []}}))
    assert resp.status_code == 200
    out = json.loads(resp.get_body())
    assert out["poe"]["status"] == "skipped" and "no queue wired" in out["poe"]["reason"]
    assert "latest.json" in out["published"]


def test_publish_estimate_no_build_poe_flag_skips_the_run(monkeypatch):
    from deliverable import functions as dfn
    import lz.functions as lzf

    monkeypatch.setattr(dfn, "_container_client", lambda: _FakeContainer())
    monkeypatch.setattr(lzf, "stage_calc_run",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("should not be called")))
    resp = dfn.publish_estimate_route(_req({"package": P.run(), "engagement": "_default_/_default_",
                                            "compute_cost": {"x": 1}}))
    assert resp.status_code == 200
    assert json.loads(resp.get_body())["poe"] is None


# --- dashboard web app ------------------------------------------------

@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setenv("FOUNDRY_PROJECT_ENDPOINT", "https://example/api/projects/x")
    import importlib
    import app as webapp
    importlib.reload(webapp)
    from fastapi.testclient import TestClient
    return webapp, TestClient(webapp.app)


def test_dashboard_page_serves_the_html(client):
    _webapp, c = client
    r = c.get("/dashboard")
    assert r.status_code == 200
    assert "Migration Assessment" in r.text
    assert "/dashboard/data" in r.text
    for fmt in ("xlsx", "docx", "pptx"):
        assert f"/dashboard/download/{fmt}" in r.text


def test_chat_page_has_new_chat_working_indicator_and_cards(client):
    _webapp, c = client
    r = c.get("/")
    assert r.status_code == 200
    assert "New chat" in r.text                       # clear / new session (markup)
    js = c.get("/static/chat.js").text                # behaviour lives in the served asset (E12.7)
    assert "The estimator is working" in js           # progress indicator
    assert "/api/prompt_cards" in js
    assert "cardgrid" in c.get("/static/chat.css").text


def test_prompt_cards_endpoint_returns_intro_and_cards(client):
    _webapp, c = client
    j = c.get("/api/prompt_cards").json()
    assert j["intro"]["title"] and j["intro"]["body"]
    assert isinstance(j["intro"]["capabilities"], list) and j["intro"]["capabilities"]
    labels = [card["label"] for card in j["cards"]]
    assert "Full estimate" in labels and "Landing zone" in labels
    assert all(card.get("prompt") for card in j["cards"])


def test_chat_page_has_engagement_picker(client):
    _webapp, c = client
    r = c.get("/")
    assert "engsel" in r.text and "New engagement" in r.text
    js = c.get("/static/chat.js").text
    assert "loadEngagements" in js and "/api/engagements" in js
    # every question rides with the engagement id — the user never types it
    assert "engagement:ENG" in js


def test_calc_regions_endpoint(client):
    _webapp, c = client
    regs = c.get("/api/calc_regions").json()["regions"]
    assert "swedencentral" in regs and "westeurope" in regs
    assert "mars-central" not in regs


def test_chat_prepends_engagement_scope(client, monkeypatch):
    webapp, c = client
    seen = {}

    class _Resp:
        id = "resp_1"
        status = "completed"
        output_text = "ok"
        output = []

    class _Responses:
        def create(self, **kw):
            seen.update(kw)
            return _Resp()

    class _OpenAI:
        responses = _Responses()

    monkeypatch.setattr(webapp, "_openai_client", lambda: _OpenAI())
    monkeypatch.setattr(webapp, "AGENT_NAME", "landfall-migration-estimator")
    # the engagement is access-checked (E8.6) — stub the guard + the stored chat
    monkeypatch.setattr(webapp, "_engagement",
                        lambda c_, p_, req=None: (f"{c_}/{p_}", f"engagements/{c_}/{p_}"))
    monkeypatch.setattr(webapp, "_load_chat", lambda eid: {})
    monkeypatch.setattr(webapp, "_save_chat", lambda eid, doc: None)
    r = c.post("/api/chat", json={"message": "how many prod servers?",
                                  "engagement": "contoso/dc-exit"})
    assert r.status_code == 200
    assert "contoso/dc-exit" in seen["input"]
    assert "engagement` argument for every tool call" in seen["input"]
    assert seen["input"].strip().endswith("how many prod servers?")


def test_dashboard_data_404_when_nothing_published(client):
    webapp, c = client
    webapp._blob_state.clear()
    from unittest.mock import patch
    with patch.object(webapp, "_read_estimate_blob", return_value=None):
        assert c.get("/dashboard/data").status_code == 404


def test_dashboard_data_returns_the_package(client):
    webapp, c = client
    pkg = P.run()
    from unittest.mock import patch
    with patch.object(webapp, "_read_estimate_blob",
                      return_value=json.dumps(pkg).encode()):
        r = c.get("/dashboard/data")
        assert r.status_code == 200
        assert r.json()["meta"]["package_id"] == pkg["meta"]["package_id"]


def test_dashboard_download_streams_blob_with_right_mime(client):
    webapp, c = client
    from unittest.mock import patch
    with patch.object(webapp, "_read_estimate_blob", return_value=b"PK\x03\x04fake"):
        r = c.get("/dashboard/download/xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        assert "attachment" in r.headers["content-disposition"]
    assert c.get("/dashboard/download/pdf").status_code == 400


def test_healthz_reports_publish_state(client):
    webapp, c = client
    from unittest.mock import patch
    with patch.object(webapp, "_read_estimate_blob", return_value=b"{}"):
        j = c.get("/healthz").json()
    assert j["ok"] and j["estimate_published"] is True
