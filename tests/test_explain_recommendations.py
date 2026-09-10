"""Cycle 63 — Unit tests for Recommendation Explanations, Provenance, and Readiness Experience (E15.4 / §E15D)."""
from __future__ import annotations

import io
import json
import os
import sys

import pytest
from conftest import ROOT

sys.path.insert(0, os.path.join(ROOT, "src", "web"))
sys.path.insert(0, os.path.join(ROOT, "evals"))

import pipeline as P  # noqa: E402
from deliverable.provenance import get_provenance_metadata  # noqa: E402
from recommendations.explain import (  # noqa: E402
    CORE_EXPLANATIONS,
    compare_baseline_to_live_guidance,
    explain_recommendation,
)


# --- 1. Recommendation Explanations (E15D.4) -----------------------------


def test_core_explanations_have_strict_six_part_anatomy():
    """Every architectural recommendation justification must satisfy the strict 6-part anatomy."""
    expected_fields = {
        "topic",
        "recommendation",
        "customer_driver",
        "landfall_rule",
        "microsoft_guidance",
        "confidence",
        "assumptions_gaps",
    }
    assert len(CORE_EXPLANATIONS) >= 6

    for key, explanation in CORE_EXPLANATIONS.items():
        d = explanation.to_dict()
        assert set(d.keys()) == expected_fields, f"Topic '{key}' missing required 6-part fields"
        for field in expected_fields:
            assert isinstance(d[field], str) and d[field].strip(), f"Field '{field}' in '{key}' is empty"
        assert d["confidence"] in ("High", "Medium", "Low")


def test_explain_recommendation_query_matching():
    """explain_recommendation resolves natural questions to the correct architectural domain."""
    firewall = explain_recommendation("Why did you recommend Azure Firewall Premium with TLS inspection?")
    assert "Firewall" in firewall["recommendation"]
    assert "PCI-DSS" in firewall["customer_driver"]
    assert firewall["confidence"] == "High"

    resiliency = explain_recommendation("What is the basis for the DR and RTO/RPO target?")
    assert "Zone-Redundant" in resiliency["recommendation"]
    assert "Tier 1" in resiliency["customer_driver"]

    sizing = explain_recommendation("Explain VM rightsizing and reserved instances")
    assert "Right-Sizing" in sizing["recommendation"]
    assert "P95" in sizing["recommendation"]

    replatform = explain_recommendation("Why replatform databases to Azure SQL and Postgres?")
    assert "Replatform" in replatform["recommendation"]

    resource = explain_recommendation("How did you size the migration team FTE count?")
    assert "FTE" in resource["recommendation"]
    assert "6R disposition" in resource["landfall_rule"]

    lz = explain_recommendation("Explain network topology and hub-spoke structure")
    assert "Hub-Spoke" in lz["recommendation"]


# --- 2. Assessment Baseline vs. Live Guidance Compare (E15D.5) -----------


def test_baseline_compare_in_sync_when_no_divergence():
    """Baseline remains authoritative and in-sync when live guidance agrees or is absent."""
    res_empty = compare_baseline_to_live_guidance(
        baseline_topic="Firewall",
        baseline_value="Azure Firewall Premium",
        live_guidance_text=None,
    )
    assert res_empty["is_divergent"] is False
    assert res_empty["flag"] == "IN_SYNC"
    assert "Authoritative" in res_empty["recommendation"] or "authoritative" in res_empty["recommendation"]

    res_matching = compare_baseline_to_live_guidance(
        baseline_topic="Firewall",
        baseline_value="Azure Firewall Premium",
        live_guidance_text="Microsoft recommends Azure Firewall Premium for Hub-Spoke enterprise deployments.",
    )
    assert res_matching["is_divergent"] is False
    assert res_matching["flag"] == "IN_SYNC"


def test_baseline_compare_flags_divergence_without_silent_mutation():
    """Live guidance differences (e.g. deprecation, preview) are flagged for architect review."""
    res_diff = compare_baseline_to_live_guidance(
        baseline_topic="App Service",
        baseline_value="Standard V2",
        live_guidance_text="Notice: App Service Standard V2 will be retired. Upgrade to V3 recommended.",
    )
    assert res_diff["is_divergent"] is True
    assert res_diff["flag"] == "DIFFERENCE_FLAGGED"
    assert res_diff["baseline_status"] == "IMMUTABLE PINNED ASSESSMENT"
    assert "completed assessment baseline will not be changed automatically" in res_diff["recommendation"]


# --- 3. Provenance Metadata (E15D.6) -------------------------------------


def test_provenance_metadata_structure_and_commit():
    """Assessment provenance carries engine version, MEG commit SHA, artifact hash, and dates."""
    meta = get_provenance_metadata()
    assert meta["engine_version"] == "2.4.0"
    assert meta["resource_model_version"] == "1.0.0"
    assert len(meta["meg_reference_commit"]) == 40  # Full Git commit SHA
    assert meta["meg_reference_commit"] == "09b269375dc7c48cee7541e4ca16faf02b3897ca"
    assert len(meta["meg_artifact_hash"]) == 64  # SHA-256
    assert meta["price_date"]
    assert meta["assessment_date"]
    assert "Microsoft Azure Migration Execution Guide" in meta["attribution"]


# --- 4. Dashboard Endpoints Integration ----------------------------------


@pytest.fixture()
def client(monkeypatch):
    from browser.serve import _Store
    import web_storage
    raw = _Store()
    raw.put("engagements/_default_/_default_/_engagement.json", b'{"visibility":"all"}')
    pkg = P.run()
    raw.put("engagements/_default_/_default_/estimate/latest.json", json.dumps(pkg).encode("utf-8"))
    monkeypatch.setattr(web_storage, "_raw_container", lambda: raw)
    monkeypatch.setattr(web_storage, "_estimate_container", lambda: raw)
    monkeypatch.setenv("FOUNDRY_PROJECT_ENDPOINT", "https://example/api/projects/x")
    import importlib
    import app as webapp
    importlib.reload(webapp)
    from fastapi.testclient import TestClient
    return webapp, TestClient(webapp.app)


def test_dashboard_explain_endpoint(client):
    _webapp, c = client
    r = c.get("/dashboard/explain?q=firewall")
    assert r.status_code == 200
    j = r.json()
    assert "Firewall" in j["recommendation"]
    assert j["confidence"] in ("High", "Medium", "Low")
    assert "landfall_rule" in j


def test_dashboard_provenance_endpoint(client):
    _webapp, c = client
    r = c.get("/dashboard/provenance")
    assert r.status_code == 200
    j = r.json()
    assert j["engine_version"] == "2.4.0"
    assert j["meg_reference_commit"] == "09b269375dc7c48cee7541e4ca16faf02b3897ca"


def test_dashboard_readiness_endpoint(client):
    _webapp, c = client
    r = c.get("/dashboard/readiness")
    assert r.status_code == 200
    j = r.json()
    assert "summary" in j
    assert "criteria" in j
    assert len(j["criteria"]) == 21
    assert all("state" in item and "evidence" in item for item in j["criteria"])


def test_dashboard_resource_plan_endpoint(client):
    _webapp, c = client
    r = c.get("/dashboard/resource-plan")
    assert r.status_code == 200
    j = r.json()
    assert "demand" in j
    assert "scenarios" in j
    assert "capacity" in j
    assert "commercial" in j
    assert j["demand"]["totals"]["total_person_days"] > 0


def test_dashboard_download_resource_plan_xlsx(client):
    _webapp, c = client
    r = c.get("/dashboard/download/resource-plan-xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    assert "resource-plan.xlsx" in r.headers["content-disposition"]

    # Verify openpyxl can load the 15-sheet workbook
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) == 15
    assert "01 Executive Summary" in wb.sheetnames
    assert "03 Role Catalogue" in wb.sheetnames
